"""奖励整理的确定性 Excel 解析、冲突审阅与导出。"""

from __future__ import annotations

import hashlib
import re
import threading
import time
import uuid
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.config import (
    REWARD_ROSTER_ELIGIBLE_STATUS,
    REWARD_ROSTER_MAX_ROWS_PER_SHEET,
    REWARD_ROSTER_MAX_SHEETS,
    REWARD_ROSTER_MAX_UPLOAD_BYTES,
    REWARD_ROSTER_SESSION_TTL_SECONDS,
)
from app.core.security import _assign_session_owner
from app.schemas.reward_roster import RewardAnalyzeRequest, RewardResolveRequest
from app.storage import group_workflow as workflow_store


_ACCOUNT_RE = re.compile(r"^\s*(\d+)\s*\(\s*(\d+)\s*\)\s*$")
_DATE_HEADER_ALIASES = ("奖励发放日期", "奖励日期", "日期", "date")
_SESSIONS: dict[str, dict] = {}
_SESSION_LOCK = threading.RLock()


def _now() -> float:
    return time.monotonic()


def _sweep_sessions() -> None:
    cutoff = _now() - REWARD_ROSTER_SESSION_TTL_SECONDS
    with _SESSION_LOCK:
        for session_id in [
            key for key, value in _SESSIONS.items() if value.get("touched_at", 0) < cutoff
        ]:
            _SESSIONS.pop(session_id, None)


def get_reward_session(session_id: str) -> dict:
    _sweep_sessions()
    with _SESSION_LOCK:
        session = _SESSIONS.get(session_id)
        if not session:
            session = workflow_store.read("imports", session_id)
            if session:
                session["content"] = workflow_store.read_blob(session_id)
                if session.get("analysis"):
                    session["analysis"]["conflict_row_ids"] = set(session["analysis"]["conflict_row_ids"])
                _SESSIONS[session_id] = session
        if not session:
            raise HTTPException(status_code=404, detail="奖励整理任务不存在或已过期，请重新上传 Excel")
        session["touched_at"] = _now()
        return session


def _display_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _persist_import(session: dict) -> None:
    from copy import deepcopy
    record = deepcopy({key: value for key, value in session.items() if key not in ("content", "touched_at")})
    if record.get("analysis"):
        record["analysis"]["conflict_row_ids"] = sorted(record["analysis"]["conflict_row_ids"])
    record["content_hash"] = workflow_store.save_blob(session["id"], session["content"])
    def save(current):
        current.clear()
        current.update(record)
    workflow_store.mutate("imports", session["id"], save, default={})


def _normalize_header(value: object) -> str:
    return re.sub(r"[\s_\-（）()：:]", "", _display_value(value).lower())


def _date_column_index(headers: list[str]) -> int | None:
    normalized_headers = [_normalize_header(header) for header in headers]
    normalized_aliases = [_normalize_header(alias) for alias in _DATE_HEADER_ALIASES]
    for index, header in enumerate(normalized_headers):
        if header in normalized_aliases:
            return index
    for index, header in enumerate(normalized_headers):
        if any(len(alias) >= 3 and alias in header for alias in normalized_aliases):
            return index
    return None


def _header_info(worksheet) -> tuple[int, list[str], list[list[str]], int, dict[str, int]]:
    if worksheet.max_row > REWARD_ROSTER_MAX_ROWS_PER_SHEET + 1:
        raise HTTPException(
            status_code=400,
            detail=f"sheet“{worksheet.title}”超过 {REWARD_ROSTER_MAX_ROWS_PER_SHEET} 行限制",
        )
    header_row = 0
    headers: list[str] = []
    samples: list[list[str]] = []
    nonempty_rows = 0
    date_column: int | None = None
    date_counts: dict[str, int] = defaultdict(int)
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=REWARD_ROSTER_MAX_ROWS_PER_SHEET + 1, values_only=True),
        start=1,
    ):
        values = [_display_value(value) for value in row]
        if not any(values):
            continue
        nonempty_rows += 1
        if not header_row:
            header_row = row_number
            headers = [value or f"未命名列 {index + 1}" for index, value in enumerate(values)]
            date_column = _date_column_index(headers)
            continue
        if len(samples) < 3:
            samples.append(values[: len(headers)])
        if date_column is not None and date_column < len(row):
            parsed_date = _parse_date(row[date_column])
            if parsed_date:
                date_counts[parsed_date] += 1
    if nonempty_rows > REWARD_ROSTER_MAX_ROWS_PER_SHEET:
        raise HTTPException(
            status_code=400,
            detail=f"sheet“{worksheet.title}”超过 {REWARD_ROSTER_MAX_ROWS_PER_SHEET} 行限制",
        )
    return header_row, headers, samples, max(0, nonempty_rows - 1), dict(date_counts)


def _open_workbook(content: bytes):
    try:
        return load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise HTTPException(status_code=400, detail="无法读取该 Excel，请确认文件未损坏且格式为 .xlsx") from exc


def create_reward_session(filename: str, content: bytes, login: dict | None) -> dict:
    if Path(filename).suffix.lower() != ".xlsx":
        raise HTTPException(status_code=400, detail="奖励整理目前仅支持 .xlsx 文件")
    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空")
    if len(content) > REWARD_ROSTER_MAX_UPLOAD_BYTES:
        limit_mb = REWARD_ROSTER_MAX_UPLOAD_BYTES // (1024 * 1024)
        raise HTTPException(status_code=413, detail=f"Excel 超过 {limit_mb} MB 上传限制")

    workbook = _open_workbook(content)
    try:
        if len(workbook.sheetnames) > REWARD_ROSTER_MAX_SHEETS:
            raise HTTPException(
                status_code=400,
                detail=f"Excel 超过 {REWARD_ROSTER_MAX_SHEETS} 个 sheet 限制",
            )
        sheets = []
        for index, name in enumerate(workbook.sheetnames):
            header_row, headers, samples, row_count, date_counts = _header_info(workbook[name])
            sheets.append({
                "name": name,
                "index": index,
                "header_row": header_row,
                "headers": headers,
                "sample_rows": samples,
                "row_count": row_count,
                "usable": bool(header_row and headers),
                "date_counts": date_counts,
            })
    finally:
        workbook.close()

    session_id = uuid.uuid4().hex
    session = {
        "id": session_id,
        "filename": Path(filename).name,
        "content": content,
        "sheets": sheets,
        "created_at": _now(),
        "touched_at": _now(),
        "analysis": None,
        "result": None,
    }
    _assign_session_owner(session, login)
    from app.services.feature_access import owner_identity
    session.update(owner_identity(login))
    _persist_import(session)
    _sweep_sessions()
    with _SESSION_LOCK:
        _SESSIONS[session_id] = session
    available_dates: dict[str, dict] = {}
    for sheet in sheets:
        for date_value, row_count in sheet["date_counts"].items():
            item = available_dates.setdefault(date_value, {"date": date_value, "row_count": 0, "sheets": []})
            item["row_count"] += row_count
            item["sheets"].append(sheet["name"])
    return {
        "session_id": session_id,
        "filename": session["filename"],
        "sheets": sheets,
        "available_dates": [available_dates[key] for key in sorted(available_dates)],
        "eligible_status": REWARD_ROSTER_ELIGIBLE_STATUS,
    }


def _parse_date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _display_value(value)
    if not text:
        return None
    normalized = text.replace("年", "-").replace("月", "-").replace("日", "")
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%m-%d-%Y"):
        try:
            return datetime.strptime(normalized, fmt).date().isoformat()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(normalized).date().isoformat()
    except ValueError:
        return None


def _parse_reward(value: object, field_label: str) -> int:
    if value is None or _display_value(value) == "":
        return 0
    if isinstance(value, bool):
        raise ValueError(f"{field_label}必须是非负整数")
    if isinstance(value, int):
        amount = value
    elif isinstance(value, float) and value.is_integer():
        amount = int(value)
    else:
        text = _display_value(value).replace(",", "")
        if text.startswith("="):
            raise ValueError(f"{field_label}是公式，请先在 Excel 中计算并粘贴为数值")
        if not re.fullmatch(r"\d+", text):
            raise ValueError(f"{field_label}必须是非负整数")
        amount = int(text)
    if amount < 0:
        raise ValueError(f"{field_label}不能为负数")
    return amount


def _parse_account(value: object) -> tuple[str, str, str]:
    text = _display_value(value)
    if text.startswith("="):
        raise ValueError("MLBBaccountid 是公式，请先在 Excel 中转为文本值")
    match = _ACCOUNT_RE.fullmatch(text)
    if not match:
        raise ValueError("MLBBaccountid 必须为 uid(zoneid)，例如 10736871(3018)")
    uid, zoneid = match.groups()
    return uid, zoneid, f"{uid}({zoneid})"


def _mapping_indexes(mapping) -> list[int]:
    return [index for index in [
        mapping.date_col,
        mapping.task_col,
        mapping.player_id_col,
        mapping.account_id_col,
        mapping.status_col,
        mapping.normal_reward_col,
        mapping.excellent_reward_col,
        mapping.other_reward_col,
    ] if index is not None]


def _conflict_id(account: str) -> str:
    digest = hashlib.sha256(account.encode("utf-8")).hexdigest()[:16]
    return f"account-{digest}"


def _build_conflicts(rows: list[dict]) -> tuple[list[dict], set[str]]:
    by_account: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        by_account[row["account"]].append(row)

    conflicts: list[dict] = []
    conflict_row_ids: set[str] = set()
    for account, account_rows in sorted(by_account.items()):
        sheets = {row["sheet"] for row in account_rows}
        same_task_groups: dict[tuple[str, str, str], list[dict]] = defaultdict(list)
        task_groups: dict[tuple[str, str], list[dict]] = defaultdict(list)
        for row in account_rows:
            task_key = row.get("reward_event_key") or row["task"]
            same_task_groups[(row["sheet"], row["date"], task_key)].append(row)
            task_groups[(row["date"], task_key)].append(row)
        reasons = []
        if len(sheets) > 1:
            reasons.append("cross_sheet_account")
        if any(len({row["sheet"] for row in group}) > 1 for group in task_groups.values()):
            reasons.append("cross_sheet_same_task")
        if any(len(group) > 1 for group in same_task_groups.values()):
            reasons.append("same_sheet_duplicate")
        if not reasons:
            continue
        ordered_rows = sorted(
            account_rows,
            key=lambda item: (item["date"], item["sheet"], item["task"], item["row_number"]),
        )
        conflicts.append({
            "conflict_id": _conflict_id(account),
            "account": account,
            "player_ids": sorted({row["player_id"] for row in account_rows if row["player_id"]}),
            "reasons": reasons,
            "rows": ordered_rows,
        })
        conflict_row_ids.update(row["row_id"] for row in account_rows)
    return conflicts, conflict_row_ids


def analyze_reward_session(session_id: str, request: RewardAnalyzeRequest, *, latest_roster=False) -> dict:
    session = get_reward_session(session_id)
    selected = list(dict.fromkeys(name.strip() for name in request.selected_sheets if name.strip()))
    if len(selected) != len(request.selected_sheets):
        raise HTTPException(status_code=400, detail="参与统计的 sheet 不能重复或为空")
    available = {sheet["name"]: sheet for sheet in session["sheets"]}
    unknown = [name for name in selected if name not in available]
    if unknown:
        raise HTTPException(status_code=400, detail=f"找不到 sheet：{'、'.join(unknown)}")
    missing_mappings = [name for name in selected if name not in request.mappings]
    if missing_mappings:
        raise HTTPException(status_code=400, detail=f"以下 sheet 尚未完成字段映射：{'、'.join(missing_mappings)}")

    reward_dates = []
    for value in request.reward_dates:
        parsed_date = _parse_date(value)
        if not parsed_date:
            raise HTTPException(status_code=400, detail=f"奖励日期格式无效：{value}")
        reward_dates.append(parsed_date)
    reward_dates = list(dict.fromkeys(reward_dates))
    if len(reward_dates) != len(request.reward_dates):
        raise HTTPException(status_code=400, detail="奖励日期不能重复")
    available_dates = {
        date_value
        for sheet in session["sheets"]
        for date_value in sheet.get("date_counts", {})
    }
    unknown_dates = [value for value in reward_dates if value not in available_dates]
    if unknown_dates:
        raise HTTPException(status_code=400, detail=f"Excel 中找不到奖励日期：{'、'.join(unknown_dates)}")
    reward_date_set = set(reward_dates)

    workbook = _open_workbook(session["content"])
    eligible_rows: list[dict] = []
    invalid_rows: list[dict] = []
    ignored = {"other_date": 0, "ineligible_status": 0, "zero_reward": 0, "blank": 0}
    try:
        for sheet_name in selected:
            info = available[sheet_name]
            if not info["usable"]:
                raise HTTPException(status_code=400, detail=f"sheet“{sheet_name}”没有可识别的表头")
            mapping = request.mappings[sheet_name]
            indexes = _mapping_indexes(mapping)
            if len(set(indexes)) != len(indexes):
                raise HTTPException(status_code=400, detail=f"sheet“{sheet_name}”的八个字段不能映射到同一列")
            if max(indexes) >= len(info["headers"]):
                raise HTTPException(status_code=400, detail=f"sheet“{sheet_name}”的字段映射超出表头范围")

            worksheet = workbook[sheet_name]
            for row_number, cells in enumerate(
                worksheet.iter_rows(min_row=info["header_row"] + 1, max_row=info["header_row"] + REWARD_ROSTER_MAX_ROWS_PER_SHEET),
                start=info["header_row"] + 1,
            ):
                values = [cell.value for cell in cells]
                if not any(_display_value(value) for value in values):
                    ignored["blank"] += 1
                    continue
                status = _display_value(values[mapping.status_col] if mapping.status_col < len(values) else None)
                if not latest_roster and status != REWARD_ROSTER_ELIGIBLE_STATUS:
                    ignored["ineligible_status"] += 1
                    continue
                date_value = values[mapping.date_col] if mapping.date_col < len(values) else None
                row_date = _parse_date(date_value)
                if not row_date:
                    invalid_rows.append({
                        "sheet": sheet_name,
                        "row_number": row_number,
                        "issue": "奖励日期无法识别",
                    })
                    continue
                if row_date not in reward_date_set:
                    ignored["other_date"] += 1
                    continue

                try:
                    normal = _parse_reward(values[mapping.normal_reward_col], "普通奖励")
                    excellent = (
                        _parse_reward(values[mapping.excellent_reward_col], "优秀额外奖励")
                        if mapping.excellent_reward_col is not None else 0
                    )
                    other = (
                        _parse_reward(values[mapping.other_reward_col], "其他奖励")
                        if mapping.other_reward_col is not None else 0
                    )
                    total = normal + excellent + other
                    if total == 0:
                        ignored["zero_reward"] += 1
                        continue
                    if latest_roster:
                        uid, zoneid, account = "", "", _display_value(values[mapping.account_id_col])
                    else:
                        uid, zoneid, account = _parse_account(values[mapping.account_id_col])
                    task = _display_value(values[mapping.task_col])
                    if not task:
                        raise ValueError("任务名称不能为空")
                    player_id = _display_value(values[mapping.player_id_col])
                    if latest_roster:
                        raw_player = values[mapping.player_id_col]
                        if not player_id or player_id.startswith("=") or (isinstance(raw_player, (int, float)) and abs(raw_player) >= 10**15):
                            raise ValueError("playerid 为空、公式或存在长数字精度风险，请恢复为原始文本")
                except (IndexError, ValueError) as exc:
                    invalid_rows.append({
                        "sheet": sheet_name,
                        "row_number": row_number,
                        "issue": str(exc),
                    })
                    continue

                row_id = f"{sheet_name}:{row_number}"
                eligible_rows.append({
                    "row_id": row_id,
                    "sheet": sheet_name,
                    "row_number": row_number,
                    "date": row_date,
                    "task": task,
                    "player_id": player_id,
                    "account": account,
                    "uid": uid,
                    "zoneid": zoneid,
                    "normal_reward": normal,
                    "excellent_reward": excellent,
                    "other_reward": other,
                    "total_reward": total,
                })
    finally:
        workbook.close()

    conflicts, conflict_row_ids = _build_conflicts(eligible_rows)
    analysis = {
        "reward_dates": reward_dates,
        "selected_sheets": selected,
        "eligible_rows": eligible_rows,
        "invalid_rows": invalid_rows,
        "conflicts": conflicts,
        "conflict_row_ids": conflict_row_ids,
        "ignored": ignored,
        "requires_live_roster": latest_roster,
    }
    session["analysis"] = analysis
    session["analysis_request"] = request.model_dump()
    session["result"] = None

    if not latest_roster and not conflicts and not invalid_rows:
        session["result"] = _build_result(analysis, eligible_rows)

    _persist_import(session)
    return _analysis_response(session)


def _analysis_response(session: dict) -> dict:
    analysis = session["analysis"]
    return {
        "session_id": session["id"],
        "reward_dates": analysis["reward_dates"],
        "selected_sheets": analysis["selected_sheets"],
        "eligible_status": REWARD_ROSTER_ELIGIBLE_STATUS,
        "eligible_row_count": len(analysis["eligible_rows"]),
        "potential_total_diamonds": sum(row["total_reward"] for row in analysis["eligible_rows"]),
        "invalid_rows": analysis["invalid_rows"],
        "conflicts": analysis["conflicts"],
        "ignored": analysis["ignored"],
        "export_ready": session["result"] is not None,
        "result": session["result"],
    }


def resolve_reward_conflicts(session_id: str, request: RewardResolveRequest) -> dict:
    session = get_reward_session(session_id)
    analysis = session.get("analysis")
    if not analysis:
        raise HTTPException(status_code=409, detail="请先完成字段映射和数据分析")
    if analysis.get("requires_live_roster"):
        raise HTTPException(409, "导入映射已保存，请在奖励整理中核对最新名单和奖励冲突，再生成正式批次")
    if analysis["invalid_rows"]:
        raise HTTPException(status_code=409, detail="Excel 中仍有无效数据，请修正后重新上传")

    conflicts = {item["conflict_id"]: item for item in analysis["conflicts"]}
    supplied = {item.conflict_id: item for item in request.resolutions}
    if set(supplied) != set(conflicts):
        missing = len(set(conflicts) - set(supplied))
        extra = len(set(supplied) - set(conflicts))
        detail = []
        if missing:
            detail.append(f"仍有 {missing} 组冲突未处理")
        if extra:
            detail.append(f"包含 {extra} 组无效冲突")
        raise HTTPException(status_code=400, detail="；".join(detail))

    selected_rows: list[dict] = [
        row for row in analysis["eligible_rows"] if row["row_id"] not in analysis["conflict_row_ids"]
    ]
    for conflict_id, conflict in conflicts.items():
        rows_by_id = {row["row_id"]: row for row in conflict["rows"]}
        selected_ids = list(dict.fromkeys(supplied[conflict_id].selected_row_ids))
        if any(row_id not in rows_by_id for row_id in selected_ids):
            raise HTTPException(status_code=400, detail=f"账号 {conflict['account']} 的处理结果包含无效记录")
        retained = [rows_by_id[row_id] for row_id in selected_ids]
        retained_tasks = [(row["date"], row["task"]) for row in retained]
        if len(retained_tasks) != len(set(retained_tasks)):
            raise HTTPException(
                status_code=400,
                detail=f"账号 {conflict['account']} 的同一任务只能保留一条奖励记录",
            )
        selected_rows.extend(retained)

    session["result"] = _build_result(analysis, selected_rows)
    _persist_import(session)
    return {
        "session_id": session_id,
        "export_ready": True,
        "result": session["result"],
    }


def _number_sort_key(value: str) -> tuple[int, int | str]:
    return (0, int(value)) if value.isdigit() else (1, value)


def _build_result(analysis: dict, selected_rows: list[dict]) -> dict:
    ordered_rows = sorted(
        selected_rows,
        key=lambda row: (row["sheet"], row["task"], row["account"], row["row_number"]),
    )
    by_account: dict[str, list[dict]] = defaultdict(list)
    by_task: dict[str, list[dict]] = defaultdict(list)
    for row in ordered_rows:
        by_account[row["account"]].append(row)
        by_task[row["task"]].append(row)

    account_totals = []
    for account, rows in by_account.items():
        account_totals.append({
            "account": account,
            "uid": rows[0]["uid"],
            "zoneid": rows[0]["zoneid"],
            "total_diamonds": sum(row["total_reward"] for row in rows),
            "source_count": len(rows),
            "player_ids": sorted({row["player_id"] for row in rows if row["player_id"]}),
            "sources": rows,
        })
    account_totals.sort(
        key=lambda item: (-item["total_diamonds"], _number_sort_key(item["zoneid"]), _number_sort_key(item["uid"])),
    )

    task_stats = []
    for task, rows in sorted(by_task.items()):
        task_stats.append({
            "task": task,
            "sheets": sorted({row["sheet"] for row in rows}),
            "dates": sorted({row["date"] for row in rows}),
            "account_count": len({row["account"] for row in rows}),
            "diamonds": sum(row["total_reward"] for row in rows),
        })

    total_diamonds = sum(item["total_diamonds"] for item in account_totals)
    simple_lines = [
        f"奖励日期：{'、'.join(analysis['reward_dates'])}",
        f"纳入 sheet：{'、'.join(analysis['selected_sheets'])}",
        f"合计：{len(account_totals)} 个账号，{total_diamonds} 钻",
        "任务统计：",
    ]
    simple_lines.extend(
        f"- {item['task']}：{item['account_count']} 人，{item['diamonds']} 钻"
        f"（来自：{'、'.join(item['sheets'])}）"
        for item in task_stats
    )
    note_simple = "\n".join(simple_lines)

    full_lines = [note_simple, "", "账号奖励明细："]
    for item in account_totals:
        full_lines.append(f"- {item['account']}：{item['total_diamonds']} 钻")
        for source in item["sources"]:
            player = f"，原始 playerid：{source['player_id']}" if source["player_id"] else ""
            full_lines.append(
                f"  - {source['sheet']}｜{source['task']}｜第 {source['row_number']} 行："
                f"{source['total_reward']} 钻{player}"
            )
    note_full = "\n".join(full_lines)

    export_accounts = sorted(
        account_totals,
        key=lambda item: (_number_sort_key(item["zoneid"]), _number_sort_key(item["uid"])),
    )
    txt = "\r\n".join(
        f"{item['zoneid']}\t{item['uid']}\t301-{item['total_diamonds']}"
        for item in export_accounts
    )
    if txt:
        txt += "\r\n"

    return {
        "unique_accounts": len(account_totals),
        "total_diamonds": total_diamonds,
        "task_stats": task_stats,
        "accounts": account_totals,
        "note_simple": note_simple,
        "note_full": note_full,
        "txt_preview": txt,
    }


def reward_download(session_id: str) -> tuple[bytes, str]:
    session = get_reward_session(session_id)
    result = session.get("result")
    if not result:
        raise HTTPException(status_code=409, detail="仍有冲突或无效数据未处理，暂不能导出")
    reward_dates = session["analysis"]["reward_dates"]
    date_label = reward_dates[0] if len(reward_dates) == 1 else f"{reward_dates[0]}等{len(reward_dates)}个日期"
    filename = f"{date_label}-奖励发放名单.txt"
    return result["txt_preview"].encode("utf-8"), filename
