"""Task lifecycle, input checking and per-response rewards. No browser execution state."""
from __future__ import annotations

from collections import defaultdict
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
import re
import uuid

from fastapi import HTTPException

from app.core import config
from app.core.parsing import parse_file_with_cell_metadata
from app.schemas.group_tasks import FeedbackMapping, TaskRules
from app.services.feature_access import owner_identity, require_permission
from app.services.group_access import record_scope, require_record
from app.storage import group_workflow as store

LABELS = ("优秀反馈", "有效反馈", "无效反馈")
ACTIVE_RUNS = {"queued", "running"}
EDITABLE_SETTINGS = {"name", "payout_date", "base_reward", "excellent_reward", "group_source_id", "group_name", "group_source_ids", "group_names"}


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def fingerprint(value) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def _event(task: dict, login: dict | None, action: str, details=None) -> None:
    who = owner_identity(login)
    task.setdefault("events", []).append({
        "id": uuid.uuid4().hex, "at": now(), "actor_key": who["owner_key"],
        "actor_name": who["owner_name"], "action": action,
        "version": task["version"], "details": deepcopy(details or {}),
    })
    from app.services.audit import _audit_log_from_login
    from app.core.runtime_logging import runtime_log
    try:
        _audit_log_from_login(None, login, "annotate", action, task["rules"].get("name", ""),
                              metadata={"task_id": task["id"], "version": task["version"]})
    except Exception:
        runtime_log("审计记录", "warning", "质量任务的全局审计记录写入失败；任务内仍保留处理记录")


def load_task(task_id: str, login: dict | None, *, write=False, feature="annotate") -> dict:
    require_permission(login, feature, write=write)
    try:
        value = store.read("tasks", task_id)
    except ValueError:
        raise HTTPException(404, "任务不存在或无权访问") from None
    return require_record(None if value and value.get("deleted_at") else value, login, write=write)


def _ensure_version(task: dict, version: int, *, allow_busy=False) -> None:
    if task.get("deleted_at"):
        raise HTTPException(404, "任务已删除")
    if task["version"] != version:
        raise HTTPException(409, "任务已更新，请刷新后核对新版本再操作")
    if not allow_busy and (task.get("run") or {}).get("status") in ACTIVE_RUNS:
        raise HTTPException(409, "本任务正在排队或标注，暂时不能修改或删除；其他任务仍可操作")


def create_task(rules: TaskRules, login: dict | None) -> dict:
    require_permission(login, "annotate", write=True)
    identifier = uuid.uuid4().hex
    task = {
        "id": identifier, **owner_identity(login), "created_at": now(), "updated_at": now(),
        "version": 1, "rules": _rules_metadata(rules), "confirmed_version": None,
        "feedback": [], "uploads": [], "mapping": None, "headers": [],
        "duplicate_decisions": {}, "player_adv": {}, "events": [], "runs": [], "run": None,
        "roster_snapshot": None, "roster_error": "", "ai_confirmation_complete": False,
    }
    _event(task, login, "创建质量标注任务")
    store.create("tasks", identifier, task)
    return task_view(task, login)


def _change(task_id: str, login: dict | None, version: int, action: str, mutate) -> dict:
    load_task(task_id, login, write=True)
    def apply(task):
        require_record(task, login, write=True)
        _ensure_version(task, version)
        details = mutate(task)
        task["version"] += 1
        task["confirmed_version"] = None
        task["updated_at"] = now()
        _refresh_checks(task)
        _event(task, login, action, details)
        return task_view(task, login)
    try:
        return store.mutate("tasks", task_id, apply)
    except FileNotFoundError:
        raise HTTPException(404, "任务不存在或无权访问") from None


def _invalidate_model(row: dict) -> None:
    if row.get("ai") or row.get("quality") or row.get("manual_quality"):
        row.setdefault("judgment_history", []).append({
            "at": now(), "ai": deepcopy(row.get("ai")), "quality": deepcopy(row.get("quality")),
            "manual_quality": deepcopy(row.get("manual_quality")),
        })
    for key in ("ai", "quality", "manual_quality", "model_error", "quality_complete", "ai_complete", "confirmed_ai", "ai_reviewed", "ai_decision"):
        row.pop(key, None)


def _rules_metadata(rules, previous=None):
    updated = rules.model_dump(mode="json")
    previous = previous or {}
    identifiers = updated.get("group_source_ids")
    if identifiers is None:
        identifiers = [updated["group_source_id"]] if updated.get("group_source_id") else []
    identifiers = list(dict.fromkeys(identifiers))
    previous_ids = previous.get("group_source_ids")
    if previous_ids is None:
        previous_ids = [previous["group_source_id"]] if previous.get("group_source_id") else []
    previous_names = dict(zip(previous_ids, previous.get("group_names") or [previous.get("group_name", "")]))
    names = []
    for identifier in identifiers:
        source = store.read("sources", identifier)
        if source and source.get("enabled"):
            names.append(source["name"])
        elif identifier in previous_ids:
            names.append(previous_names.get(identifier, identifier))
        else:
            raise HTTPException(422, "任务归属群组已停用或不存在，请重新选择")
    updated.update(group_source_ids=identifiers, group_names=names,
                   group_source_id=next(iter(identifiers), ""), group_name="、".join(names))
    return updated


def update_rules(task_id, version, rules: TaskRules, login):
    def mutate(task):
        previous = task["rules"]
        updated = _rules_metadata(rules, previous)
        if rules_locked(task) and any(previous.get(key) != value for key, value in updated.items() if key not in EDITABLE_SETTINGS):
            raise HTTPException(409, "新建配置已完成，数据校验、重复提交和标注依据不能修改；请新建任务")
        task["rules"] = updated
        task["rules_initialized"] = not rule_issues(task)
        return {"before": previous, "after": updated, "model_context_changed": False}
    return _change(task_id, login, version, "修改任务规则", mutate)


def rules_locked(task):
    return bool(task.get("rules_initialized") or task.get("uploads") or not rule_issues(task))


def annotation_inputs_locked(task):
    return bool(task.get("run") or task.get("runs") or any(row.get("ai") or row.get("quality") for row in task.get("feedback", [])))


def update_settings(task_id, request, login):
    current = load_task(task_id, login, write=True)
    values = request.model_dump(mode="json", exclude={"expected_version"}, exclude_unset=True)
    if "group_source_id" in values and "group_source_ids" not in values:
        values["group_source_ids"] = None  # Legacy callers still send a single group.
    if not values["name"].strip():
        raise HTTPException(422, "请填写任务名称")
    return update_rules(task_id, request.expected_version, TaskRules(**{**current["rules"], **values}), login)


def delete_task(task_id, version, login):
    load_task(task_id, login, write=True)
    def validate(task):
        require_record(task, login, write=True)
        _ensure_version(task, version)
    try:
        store.archive_task(task_id, validate, {"deleted_at": now(), "deleted_by": owner_identity(login)})
    except FileNotFoundError:
        raise HTTPException(409, "任务或原始文件缺失，无法建立完整恢复备份；任务未删除") from None
    except (ValueError, OSError):
        raise HTTPException(409, "恢复备份校验未通过，任务未删除，请联系管理员检查存储") from None
    return {"id": task_id, "deleted": True}


def _mapping_guess(headers: list[str]) -> dict:
    def find(words):
        return next((i for i, h in enumerate(headers) if re.sub(r"[ _-]", "", str(h)).casefold() in words), None)
    player = find({"playerid", "discordid", "whatsappid", "玩家id", "discord/whatsappid"})
    count = find({"实际对局次数", "实际体验次数", "对局次数", "matches", "matchcount", "actualmatches"})
    adv = find({"adv", "adv账号", "advaccount"})
    return {"id_col": player, "count_col": count, "adv_col": adv, "open_text_cols": []}


def upload_feedback(task_id, version, filename: str, content: bytes, login):
    current = load_task(task_id, login, write=True)
    _ensure_version(current, version)
    if current.get("uploads"):
        raise HTTPException(409, "该任务已上传反馈，不能替换原始文件；反馈变动请新建任务")
    if rule_issues(current):
        raise HTTPException(409, {"message": "请先完成任务规则，再上传反馈", "issues": rule_issues(current)})
    if len(content) > config.REWARD_ROSTER_MAX_UPLOAD_BYTES:
        raise HTTPException(413, "上传反馈超过大小限制")
    try:
        rows, cell_issues = parse_file_with_cell_metadata(filename, content)
    except (ValueError, OSError) as exc:
        raise HTTPException(400, str(exc)) from exc
    if len(rows) < 2:
        raise HTTPException(422, "文件为空或只有表头")
    if len(rows) - 1 > config.GROUP_FEEDBACK_MAX_ROWS:
        raise HTTPException(413, f"反馈数量超过配置上限 {config.GROUP_FEEDBACK_MAX_ROWS}")
    upload_id = uuid.uuid4().hex
    def mutate(task):
        if task.get("uploads"):
            raise HTTPException(409, "该任务已上传反馈，反馈变动请新建任务")
        digest = store.save_blob(upload_id, content)
        if task["feedback"]:
            task.setdefault("feedback_versions", []).append({
                "version": task["version"], "headers": task["headers"], "mapping": task["mapping"],
                "feedback": task["feedback"],
            })
        task["headers"] = [str(item) if item is not None else "" for item in rows[0]]
        task["mapping"] = None
        task["mapping_suggestion"] = _mapping_guess(task["headers"])
        from app.services.annotate_workflow import group_column_metadata
        task.update(group_column_metadata([task["headers"], *rows[1:]]))
        task["roster_snapshot"] = None
        task["roster_error"] = ""
        task["feedback"] = [
            {"id": uuid.uuid4().hex, "source_row": index + 1, "upload_id": upload_id,
             "cells": [str(value) if value is not None else "" for value in row],
             "cell_issues": {str(c): cell_issues[f"{index}:{c}"] for c in range(len(row)) if f"{index}:{c}" in cell_issues},
             "corrections": {}, "reviews": []}
            for index, row in enumerate(rows[1:], 1) if any(str(value).strip() for value in row)
        ]
        if not task["feedback"]:
            raise HTTPException(422, "没有非空反馈行")
        task["uploads"].append({"id": upload_id, "filename": str(filename).replace("\\", "/").split("/")[-1],
                                "sha256": digest, "at": now()})
        task["duplicate_decisions"] = {}
        task["player_adv"] = {}
        task["ai_confirmation_complete"] = False
        return {"upload_id": upload_id, "rows": len(task["feedback"])}
    return _change(task_id, login, version, "上传反馈", mutate)


async def translate_upload_headers(task_id, version, login):
    from app.services.annotate_workflow import _translate_headers
    current = load_task(task_id, login, write=True)
    _ensure_version(current, version)
    translated, warning = await _translate_headers(current["headers"])
    def save(task):
        require_record(task, login, write=True)
        _ensure_version(task, version)
        task.update(headers_zh=translated, header_translation_warning=warning)
        return task_view(task, login)
    return store.mutate("tasks", task_id, save)


def set_mapping(task_id, version, mapping: FeedbackMapping, login, annotation_options=None):
    def mutate(task):
        if rule_issues(task) or not task["feedback"]:
            raise HTTPException(409, "请先完成任务规则并上传反馈")
        values = mapping.model_dump()
        if task["rules"]["requires_count"] and values["count_col"] is None:
            raise HTTPException(422, "本任务要求校验次数，请选择数据校验列")
        indexes = [values["id_col"], *values["open_text_cols"]]
        indexes.extend(values[key] for key in ("count_col", "adv_col") if values[key] is not None)
        if len(indexes) != len(set(indexes)) or any(i < 0 or i >= len(task["headers"]) for i in indexes):
            raise HTTPException(422, "身份、回答、次数和 ADV 列不能重叠，且必须在表头范围内")
        options = annotation_options.model_dump() if annotation_options is not None else None
        if options is not None and not (options["ai_detect"] or options["quality"]):
            raise HTTPException(422, "请至少选择一项标注任务")
        options_changed = options is not None and options != annotation_settings(task)
        if annotation_inputs_locked(task) and (task["mapping"] != values or options_changed):
            raise HTTPException(409, "标注已开始，题目列、标注任务和调研背景不能修改；请新建任务")
        if task["mapping"] != values or options_changed:
            for row in task["feedback"]:
                _invalidate_model(row)
                row["corrections"] = {}
            task["duplicate_decisions"] = {}
            task["ai_confirmation_complete"] = False
            task["roster_snapshot"] = None
            task["roster_error"] = ""
        task["mapping"] = values
        if options is not None:
            task["annotation_options"] = options
            task["rules"]["ai_detect"] = options["ai_detect"]
        return values
    return _change(task_id, login, version, "确认上传列映射", mutate)


def _cell(row, column):
    return str(row["cells"][column]).strip() if column is not None and column < len(row["cells"]) else ""


def _row(task, response_id):
    row = next((item for item in task["feedback"] if item["id"] == response_id), None)
    if not row:
        raise HTTPException(404, "反馈记录不存在")
    return row


def _parse_count(value):
    text = str(value).strip()
    if not re.fullmatch(r"[0-9]+(?:\.0+)?", text):
        return None
    number = int(text.split(".")[0])
    return number if number <= 2147483647 else None


def _refresh_checks(task: dict) -> None:
    mapping = task.get("mapping")
    task["input_issues"] = []
    task["duplicate_groups"] = []
    if not mapping:
        if task["feedback"]:
            task["input_issues"].append({"kind": "mapping", "message": "请确认身份列和完整反馈题目列"})
        return
    snapshot = task.get("roster_snapshot")
    entries = (snapshot or {}).get("entries", [])
    by_player = defaultdict(list)
    for entry in entries:
        by_player[entry["player_id"]].append(entry)
    if not snapshot or task.get("roster_error"):
        task["input_issues"].append({"kind": "roster", "message": task.get("roster_error") or "请读取已配置的飞书名单以核查身份"})
    for row in task["feedback"]:
        edits = row["corrections"]
        row["player_id"] = str(edits.get("player_id", _cell(row, mapping["id_col"]))).strip()
        row["adv"] = str(edits.get("adv", _cell(row, mapping.get("adv_col")))).strip()
        row["excluded"] = False
        row["issues"] = []
        if not row["player_id"]:
            row["issues"].append({"kind": "identity", "message": "playerid 为空，请补填原始身份"})
        elif str(mapping["id_col"]) in row.get("cell_issues", {}) and "player_id" not in edits:
            row["issues"].append({"kind": "identity", "message": row["cell_issues"][str(mapping["id_col"])]})
        if row["adv"] and str(mapping.get("adv_col")) in row.get("cell_issues", {}) and "adv" not in edits:
            row["issues"].append({"kind": "adv", "message": "ADV：" + row["cell_issues"][str(mapping["adv_col"])]})
        candidates = by_player.get(row["player_id"], [])
        chosen = [item for item in candidates if item["key"] == edits.get("candidate_key")]
        row["identity_candidates"] = deepcopy(candidates)
        row["identity"] = deepcopy(chosen[0] if len(chosen) == 1 else candidates[0] if len(candidates) == 1 else None)
        if snapshot and row["player_id"] and not row["identity"]:
            row["issues"].append({"kind": "roster_identity", "message": "身份匹配多个名单条目，请明确选择" if candidates else "名单中找不到 playerid，请修正身份或维护名单后重新读取"})
        if row["identity"] and row["identity"].get("identity_error"):
            row["issues"].append({"kind": "roster_identity", "message": row["identity"]["identity_error"]})
        raw_count = edits.get("count", _cell(row, mapping.get("count_col")))
        count = _parse_count(raw_count)
        row["count"] = count
        if not task["rules"]["requires_count"]:
            row["gate"] = {"status": "not_applicable", "reason": "本任务不要求数据校验"}
        elif count is None:
            row["gate"] = {"status": "missing", "reason": "实际次数为空或无法解析，请补填；不会按 0 处理"}
            row["issues"].append({"kind": "count", "message": row["gate"]["reason"]})
        elif not task["rules"].get("min_count"):
            row["gate"] = {"status": "missing", "reason": "最低次数尚未配置"}
        elif count < task["rules"]["min_count"]:
            row["gate"] = {"status": "fail", "reason": f"实际完成 {count} 场，要求至少 {task['rules']['min_count']} 场；按任务规则判为无效，未进行 AI 判断。"}
        else:
            row["gate"] = {"status": "pass", "reason": f"实际完成 {count} 场，达到至少 {task['rules']['min_count']} 场要求"}
    if not task["rules"]["allow_repeat"]:
        def check_group(members, *, roster_only):
            if len(members) < 2:
                return
            ids = sorted(item["id"] for item in members)
            key = fingerprint(ids)[:24]
            retained = task["duplicate_decisions"].get(key)
            resolved = retained in ids
            task["duplicate_groups"].append({"id": key, "response_ids": ids, "keep_response_id": retained if resolved else None, "roster_only": roster_only})
            for row in members:
                row["excluded"] = resolved and row["id"] != retained
            if not resolved:
                task["input_issues"].append({"kind": "roster_duplicate" if roster_only else "duplicate", "group_id": key, "message": "身份重复，请查看完整反馈并选择保留一份"})
        # Keep upload duplicate choices stable when a later roster read reveals
        # that different player IDs share one payout account.
        player_groups = defaultdict(list)
        for row in task["feedback"]:
            if row["player_id"]:
                player_groups[row["player_id"]].append(row)
        for members in player_groups.values():
            check_group(members, roster_only=False)
        account_groups = defaultdict(list)
        for row in task["feedback"]:
            ori = (row.get("identity") or {}).get("ori")
            if ori and not row["excluded"]:
                account_groups[ori].append(row)
        for members in account_groups.values():
            if len({row["player_id"] for row in members}) > 1:
                check_group(members, roster_only=True)
    for row in task["feedback"]:
        if not row["excluded"]:
            task["input_issues"].extend({**issue, "response_id": row["id"], "source_row": row["source_row"]} for issue in row["issues"])


async def refresh_roster(task_id, version, login):
    from app.services.roster_sources import read_roster_snapshot
    task = load_task(task_id, login, write=True)
    _ensure_version(task, version)
    _refresh_checks(task)
    if readiness_issues(task) or not annotation_complete(task):
        raise HTTPException(409, "请先完成质量标注和 AI 作答确认，再核对玩家名单")
    try:
        snapshot = await read_roster_snapshot(login)
    except Exception as exc:
        def fail(current):
            require_record(current, login, write=True)
            _ensure_version(current, version)
            current["roster_error"] = str(getattr(exc, "detail", exc))
            _refresh_checks(current)
            return task_view(current, login)
        result = store.mutate("tasks", task_id, fail)
        return {"ok": False, "task": result}
    def mutate(current):
        current["roster_snapshot"] = snapshot
        current["roster_error"] = ""
        return {"snapshot_id": snapshot["id"], "read_at": snapshot["read_at"]}
    return {"ok": True, "task": _change(task_id, login, version, "刷新名单并核查身份", mutate)}


def correct_feedback(task_id, response_id, request, login):
    def mutate(task):
        row = _row(task, response_id)
        updates = request.model_dump(exclude_unset=True, exclude={"expected_version", "note"})
        if {"player_id", "count"} & updates.keys():
            raise HTTPException(409, "原始反馈和数据校验值不能修改；请修正文件后新建任务")
        updates = {key: value for key, value in updates.items() if value is not None or key == "count"}
        previous = deepcopy(row["corrections"])
        row["corrections"].update(updates)
        task["duplicate_decisions"] = {}
        if "player_id" in updates and "candidate_key" not in updates and previous.get("player_id") != updates["player_id"]:
            row["corrections"].pop("candidate_key", None)
        row["reviews"].append({"kind": "input", "at": now(), **owner_identity(login),
                               "before": previous, "after": deepcopy(row["corrections"]), "note": request.note.strip()})
        return {"response_id": response_id, "changes": updates, "note": request.note.strip()}
    return _change(task_id, login, request.expected_version, "修正反馈输入", mutate)


def resolve_duplicate(task_id, request, login):
    def mutate(task):
        _refresh_checks(task)
        group = next((item for item in task["duplicate_groups"] if item["id"] == request.group_id), None)
        if not group or request.keep_response_id not in group["response_ids"]:
            raise HTTPException(409, "重复组已变化，请重新核对")
        task["duplicate_decisions"][request.group_id] = request.keep_response_id
        return {"group_id": request.group_id, "keep_response_id": request.keep_response_id}
    return _change(task_id, login, request.expected_version, "处理重复反馈", mutate)


def quality_result_complete(task: dict, row: dict) -> bool:
    """A cached done flag cannot stand in for labels on the selected questions."""
    import annotate
    result = row.get("quality") or {}
    columns = (task.get("mapping") or {}).get("open_text_cols") or []
    labels = result.get("q_labels") or {}
    return bool(row.get("quality_complete") and columns
                and annotate.canonical_quality_label(result.get("overall"), overall=True) in LABELS
                and all(not _cell(row, col) or annotate.canonical_quality_label(labels.get(f"col_{col}")) in LABELS
                        for col in columns))


def effective_result(task: dict, row: dict) -> dict:
    if row.get("excluded"):
        return {"complete": True, "source": "excluded", "label": None, "reason": "重复记录已排除，保留原文供追溯", "base": 0, "excellent": 0}
    gate = row.get("gate") or {}
    label, reason, source = None, "", "pending"
    if gate.get("status") == "fail":
        label, reason, source = "无效反馈", gate["reason"], "rule"
    elif row.get("confirmed_ai"):
        label, reason, source = "无效反馈", "操作者已确认 AI 作答，按无效反馈处理", "human_ai"
    elif not any(issue["kind"] not in ("roster_identity", "adv") for issue in row.get("issues", [])) and gate.get("status") in ("pass", "not_applicable"):
        if quality_result_complete(task, row) and (not task["rules"]["ai_detect"] or row.get("ai_complete")):
            judgment = row.get("manual_quality") or row.get("quality") or {}
            label = judgment.get("label") or judgment.get("overall")
            reason = judgment.get("note") or judgment.get("overall_reason") or ""
            source = "human" if row.get("manual_quality") else "ai"
    if label not in LABELS:
        return {"complete": False, "source": "pending", "label": None,
                "reason": row.get("model_error") or "输入或质量结果尚未完整，可处理异常或重试缺项", "base": None, "excellent": None}
    base = (task["rules"].get("base_reward") or 0) if label != "无效反馈" else 0
    extra = (task["rules"].get("excellent_reward") or 0) if label == "优秀反馈" else 0
    return {"complete": True, "source": source, "label": label, "reason": reason, "base": base, "excellent": extra}


def override_quality(task_id, response_id, request, login):
    def mutate(task):
        row = _row(task, response_id)
        if row.get("excluded") or row.get("confirmed_ai") or (row.get("gate") or {}).get("status") == "fail":
            raise HTTPException(409, "请先处理排除、AI 作答确认或数据门槛；不能用质量标签绕过规则")
        if not quality_result_complete(task, row):
            raise HTTPException(409, "质量结果尚未完整，请先补齐")
        if request.column_index is not None:
            return _review_question(task, row, request, login)
        before = deepcopy(row.get("manual_quality"))
        if request.restore:
            row.pop("manual_quality", None)
        else:
            if request.label not in LABELS:
                raise HTTPException(422, "请选择整体质量")
            row["manual_quality"] = {"label": request.label, "note": request.note.strip(), "at": now(), **owner_identity(login)}
        row["reviews"].append({"kind": "quality", "at": now(), **owner_identity(login),
                               "before": before, "after": deepcopy(row.get("manual_quality")), "note": request.note.strip()})
        return {"response_id": response_id, "before": before, "after": row.get("manual_quality"), "note": request.note.strip()}
    return _change(task_id, login, request.expected_version, "人工调整整体质量", mutate)


def _review_question(task, row, request, login):
    """The historical per-question interaction, retaining independent overall judgments."""
    import annotate
    column = request.column_index
    if column not in task["mapping"]["open_text_cols"] or not _cell(row, column):
        raise HTTPException(422, "只能修改已作答的主观题；未作答题固定为 N/A")
    result = row["quality"]
    key = f"col_{column}"
    labels, reasons = result.setdefault("q_labels", {}), result.setdefault("q_reasons", {})
    original = labels.get(key)
    if original not in LABELS:
        raise HTTPException(409, "当前题标签不完整，请先补齐质量打标")
    baseline = result.setdefault("quality_review_baseline", {}).setdefault(key, {"label": original, "reason": reasons.get(key, "")})
    label = baseline["label"] if request.restore else request.label
    if label not in LABELS:
        raise HTTPException(422, "请选择有效的逐题质量标签")
    reviews = result.setdefault("human_reviews", {})
    labels[key] = label
    if label == baseline["label"]:
        reasons[key] = baseline["reason"]
        reviews.pop(key, None)
    else:
        reasons[key] = f"人工复核调整：{baseline['label']} → {label}；AI 原判断：{baseline['reason'] or '未提供判断依据'}"
        reviews[key] = {"from_label": baseline["label"], "to_label": label, "reviewed_at": now(), **owner_identity(login)}
    from app.services.annotate_workflow import _is_holistic_quality
    if _is_holistic_quality(result):
        if reviews:
            result["overall_review_note"] = f"人工复核调整{len(reviews)}道题；整体质量保留模型对完整作答的原判断，未重新评估"
        else:
            result.pop("overall_review_note", None)
    else:
        signals = annotate.detect_low_effort_signals(row["cells"], task["headers"], task["mapping"]["open_text_cols"], task["mapping"]["id_col"], labels, headers_zh=task.get("headers_zh", []))
        result["overall"], result["overall_reason"] = annotate.calculate_overall_quality(labels, task["mapping"]["open_text_cols"], low_effort=signals)
    row["reviews"].append({"kind": "question_quality", "column_index": column, "before": original, "after": label, "at": now(), **owner_identity(login)})
    return {"response_id": row["id"], "column_index": column, "before": original, "after": label}


def review_ai(task_id, response_id, request, login):
    def mutate(task):
        state = annotation_status(task)
        if not state["ai_complete"] or response_id not in state["review_ids"]:
            raise HTTPException(409, "请先补齐 AI 识别结果，再复核候选玩家")
        row = _row(task, response_id)
        row["ai_reviewed"], row["ai_decision"] = True, request.is_ai
        task["ai_confirmation_complete"] = False
        return {"response_id": response_id, "is_ai": request.is_ai}
    return _change(task_id, login, request.expected_version, "人工复核 AI 作答", mutate)


def confirm_ai(task_id, request, login):
    def mutate(task):
        if readiness_issues(task) or not task["rules"]["ai_detect"]:
            raise HTTPException(409, "请先完成输入检查，并开启 AI 作答识别")
        eligible = [row for row in task["feedback"] if not row.get("excluded") and row.get("gate", {}).get("status") in ("pass", "not_applicable")]
        ids = {row["id"] for row in eligible}
        selected = set(request.response_ids)
        if not selected <= ids or any(not row.get("ai_complete") for row in eligible):
            raise HTTPException(409, "AI 作答识别尚有缺项或选择已失效，请先补齐")
        if task.get("annotation_options"):
            review_ids = set(annotation_status(task)["review_ids"])
            if any(not row.get("ai_reviewed") for row in eligible if row["id"] in review_ids):
                raise HTTPException(409, "请逐位完成人工确认")
            if selected != {row["id"] for row in eligible if row["id"] in review_ids and row.get("ai_decision")}:
                raise HTTPException(409, "选择与已保存的人工复核不一致，请刷新后核对")
        for row in eligible:
            before = bool(row.get("confirmed_ai"))
            row["confirmed_ai"] = row["id"] in selected
            if before != row["confirmed_ai"]:
                row["reviews"].append({"kind": "ai_confirmation", "at": now(), **owner_identity(login),
                                       "before": before, "after": row["confirmed_ai"]})
        task["ai_confirmation_complete"] = True
        return {"confirmed_response_ids": sorted(selected)}
    return _change(task_id, login, request.expected_version, "人工确认 AI 作答", mutate)


def set_player_adv(task_id, request, login):
    def mutate(task):
        if request.player_id not in {row.get("player_id") for row in task["feedback"]}:
            raise HTTPException(404, "任务中不存在此玩家")
        task["player_adv"][request.player_id] = {"value": request.adv.strip(), "note": request.note.strip(),
                                                "at": now(), **owner_identity(login)}
        return {"player_id": request.player_id, "adv": request.adv.strip(), "note": request.note.strip()}
    return _change(task_id, login, request.expected_version, "确认玩家 ADV", mutate)


def rule_issues(task: dict) -> list[dict]:
    rules = task["rules"]
    issues = []
    if not rules["name"].strip() or not rules["payout_date"] or rules["base_reward"] is None or rules["excellent_reward"] is None:
        issues.append({"kind": "rules", "message": "请确认任务名称、发奖日期和两项奖励金额（允许 0）"})
    if rules["requires_count"] and not rules.get("min_count"):
        issues.append({"kind": "rules", "message": "请配置实际次数门槛"})
    return issues


def annotation_complete(task: dict) -> bool:
    retained = [row for row in task["feedback"] if not row.get("excluded")]
    needs_ai = task["rules"]["ai_detect"] and any(row.get("gate", {}).get("status") in ("pass", "not_applicable") for row in retained)
    return bool(retained) and all(effective_result(task, row)["complete"] for row in retained) and (not needs_ai or task.get("ai_confirmation_complete", False))


def annotation_settings(task):
    return deepcopy(task.get("annotation_options") or {"ai_detect": task["rules"]["ai_detect"], "quality": True, "background": ""})


def reward_source_task(task):
    """Payout matching belongs to the batch; ignore earlier per-task roster caches."""
    current = deepcopy(task)
    current["roster_snapshot"] = None
    current["roster_error"] = ""
    _refresh_checks(current)
    return current


def reward_ready(task):
    current = reward_source_task(task)
    return (not current.get("deleted_at") and (current.get("run") or {}).get("status") not in ACTIVE_RUNS
            and annotation_settings(current)["quality"] and not readiness_issues(current) and annotation_complete(current))


def annotation_status(task):
    options = annotation_settings(task)
    eligible = [row for row in task["feedback"] if not row.get("excluded") and row.get("gate", {}).get("status") in ("pass", "not_applicable")]
    ai_complete = all(row.get("ai_complete") for row in eligible)
    review_ids = [row["id"] for row in eligible if row.get("ai_complete") and float((row.get("ai") or {}).get("ai_prob", 0)) >= config.ANNOTATE_AI_REVIEW_THRESHOLD]
    reviewed_ids = [row["id"] for row in eligible if row["id"] in review_ids and (row.get("ai_reviewed") or task.get("ai_confirmation_complete"))]
    quality_complete = all(row.get("confirmed_ai") or quality_result_complete(task, row) for row in eligible)
    ai_confirmed = not eligible or task.get("ai_confirmation_complete", False)
    complete = not readiness_issues(task) and (not options["ai_detect"] or ai_complete and ai_confirmed) and (not options["quality"] or quality_complete)
    phase = "columns"
    if not readiness_issues(task):
        if options["ai_detect"] and not ai_complete:
            phase = "ai_progress"
        elif options["ai_detect"] and not ai_confirmed:
            phase = "ai_review"
        elif options["quality"] and not quality_complete:
            phase = "quality_progress"
        else:
            phase = "results"
    return {"options": options, "phase": phase, "complete": bool(complete), "ai_complete": ai_complete,
            "quality_complete": quality_complete, "review_ids": review_ids, "reviewed_ids": reviewed_ids,
            "review_threshold": config.ANNOTATE_AI_REVIEW_THRESHOLD, "high_threshold": config.ANNOTATE_AI_HIGH_THRESHOLD,
            "total": len(eligible), "ai_done": sum(bool(row.get("ai_complete")) for row in eligible),
            "quality_total": sum(not row.get("confirmed_ai") for row in eligible),
            "quality_done": sum(quality_result_complete(task, row) for row in eligible if not row.get("confirmed_ai"))}


def annotation_download(task_id, login):
    """Reuse the annotation exporter with stable response IDs (including repeated players)."""
    import annotate
    task = load_task(task_id, login)
    _refresh_checks(task)
    status = annotation_status(task)
    if not status["complete"]:
        raise HTTPException(409, "结果不完整，下载已被阻断。请返回对应任务重试。")
    headers = task["headers"]
    rows = [headers]
    ai, quality, confirmed = [], [], set()
    # An additional lookup cell is not part of headers, so it is never exported.
    for row in task["feedback"]:
        if row.get("excluded"):
            continue
        cells = list(row["cells"][:len(headers)])
        cells.extend([""] * (len(headers) - len(cells)))
        rows.append([*cells, row["id"]])
        if row.get("ai"):
            ai.append({**deepcopy(row["ai"]), "id": row["id"]})
        result = deepcopy(row.get("quality") or {})
        if row.get("manual_quality"):
            result.update(overall=row["manual_quality"]["label"], overall_reason=row["manual_quality"]["note"])
        if row.get("gate", {}).get("status") == "fail":
            result.update(overall="无效反馈", overall_reason=row["gate"]["reason"], q_labels={f"col_{c}": "无效反馈" if _cell(row,c) else "N/A" for c in task["mapping"]["open_text_cols"]})
        quality.append({**result, "id": row["id"]})
        if row.get("confirmed_ai"):
            confirmed.add(row["id"])
    content = annotate.generate_annotated_excel(rows, headers, ai, confirmed, quality, task["mapping"]["open_text_cols"], len(headers), status["options"])
    skipped = {index: row for index, row in enumerate((r for r in task["feedback"] if not r.get("excluded")), 2)
               if row.get("gate", {}).get("status") == "fail"}
    if skipped:
        # The legacy exporter has no count-gate concept. Keep originals and explicitly
        # distinguish rule-skipped cells from missing model output; never invent scores.
        from io import BytesIO
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content))
        sheet = book.active
        columns = {cell.value: cell.column for cell in sheet[1]}
        for index, row in skipped.items():
            reason = row["gate"]["reason"]
            if status["options"]["ai_detect"]:
                sheet.cell(index, columns["AI作答标签"], "未执行（数据校验未通过）")
                sheet.cell(index, columns["AI判断原因"], reason)
                for name in ("AI内容生成概率", "AI润色概率", "AI原文证据", "AI反向证据"):
                    sheet.cell(index, columns[name]).value = None
            for column in task["mapping"]["open_text_cols"]:
                if not _cell(row, column):
                    continue
                title = headers[column]
                sheet.cell(index, columns[f"[{title}]中文翻译"], "未执行（数据校验未通过）")
                if status["options"]["quality"]:
                    sheet.cell(index, columns[f"[{title}]质量原因"], reason)
        output = BytesIO()
        book.save(output)
        book.close()
        content = output.getvalue()
    from app.services.annotate_workflow import _annotate_download_filename
    return content, _annotate_download_filename(task["uploads"][-1]["filename"])


def stage_complete(task: dict, row: dict, stage: str) -> bool:
    if row.get("excluded") or row.get("gate", {}).get("status") == "fail":
        return True
    return bool(row.get("ai_complete")) if stage == "ai" else effective_result(task, row)["complete"]


def workflow_state(task: dict) -> dict:
    maximum = 1
    if not rule_issues(task):
        maximum = 2
        if not readiness_issues(task):
            maximum = 3
            eligible = [row for row in task["feedback"] if not row.get("excluded") and row.get("gate", {}).get("status") in ("pass", "not_applicable")]
            ai_reviewed = task.get("ai_confirmation_complete") and all(row.get("ai_complete") for row in eligible)
            if not task["rules"]["ai_detect"] or not eligible or ai_reviewed:
                maximum = 4
            if annotation_complete(task):
                maximum = 5
                if task.get("confirmed_version") == task["version"] and not readiness_issues(task, settlement=True):
                    maximum = 6
    return {"max_step": maximum, "ai_enabled": task["rules"]["ai_detect"], "annotation_complete": annotation_complete(task)}


def readiness_issues(task: dict, *, settlement=False) -> list[dict]:
    rules = task["rules"]
    issues = rule_issues(task) + deepcopy([issue for issue in task.get("input_issues", [])
                                       if settlement or issue["kind"] not in ("roster", "roster_identity", "roster_duplicate", "adv")])
    if not task["feedback"] or not task["mapping"]:
        issues.append({"kind": "upload", "message": "请上传反馈并确认完整列映射"})
    if settlement:
        retained = [row for row in task["feedback"] if not row.get("excluded")]
        if any(not effective_result(task, row)["complete"] for row in retained):
            issues.append({"kind": "incomplete", "message": "仍有未完成的反馈，不能确认结算"})
        if rules["ai_detect"] and any(row.get("gate", {}).get("status") in ("pass", "not_applicable") for row in retained) and not task["ai_confirmation_complete"]:
            issues.append({"kind": "ai_review", "message": "请核对 AI 作答识别结果，可确认没有 AI 作答"})
        from app.services.task_records import aggregate_players
        for player in aggregate_players(task):
            for issue in player["issues"]:
                issues.append({"kind": "player", "player_id": player["player_id"], "message": issue})
    return issues


def confirm_task(task_id, version, login):
    load_task(task_id, login, write=True)
    require_permission(login, "annotate", write=True)
    def apply(task):
        require_record(task, login, write=True)
        _ensure_version(task, version)
        _refresh_checks(task)
        issues = readiness_issues(task, settlement=True)
        if issues:
            raise HTTPException(409, {"message": "任务尚不能确认", "issues": issues})
        task["confirmed_version"] = task["version"]
        task["confirmed_at"] = now()
        _event(task, login, "确认任务结果")
        return task_view(task, login)
    return store.mutate("tasks", task_id, apply)


def task_view(task: dict, login, *, summary=False, access=None) -> dict:
    from app.services.task_records import aggregate_players
    # Re-evaluate legacy saved checks under the current workflow without rewriting history.
    task = deepcopy(task)
    _refresh_checks(task)
    result = deepcopy(task)
    if not summary and task.get("headers") and "column_types" not in task:
        from app.services.annotate_workflow import group_column_metadata
        metadata = group_column_metadata([task["headers"], *[row["cells"] for row in task["feedback"]]])
        result.update({key: value for key, value in metadata.items() if key not in result})
    result["can_edit"] = (access or record_scope(login)).can_read(task)
    from app.services.auth import _get_user_permission_levels
    result["can_edit"] = result["can_edit"] and _get_user_permission_levels(login).get("annotate") == "use"
    result["rules_locked"] = rules_locked(task)
    result["annotation_inputs_locked"] = annotation_inputs_locked(task)
    result["reward_ready"] = reward_ready(task)
    for saved_run in [result.get("run"), *result.get("runs", [])]:
        if saved_run:
            saved_run.pop("owner_login", None)
    retained = [row for row in task["feedback"] if not row.get("excluded")]
    outcomes = [effective_result(task, row) for row in retained]
    result["summary"] = {
        "feedback_count": len(retained), "player_count": len({row.get("player_id") for row in retained}),
        "complete_count": sum(item["complete"] for item in outcomes),
        "pending_count": sum(not item["complete"] for item in outcomes),
        "base": sum(item["base"] or 0 for item in outcomes),
        "excellent": sum(item["excellent"] or 0 for item in outcomes),
    }
    run = task.get("run") or {}
    if run.get("status") in ACTIVE_RUNS:
        status = run["status"]
    elif not task["feedback"]:
        status = "draft"
    elif readiness_issues(task):
        status = "input_issues"
    elif not annotation_settings(task)["quality"] and annotation_status(task)["complete"]:
        status = "annotation_done"
    elif task["rules"]["ai_detect"] and not task.get("ai_confirmation_complete") and all(stage_complete(task, row, "ai") for row in retained) and any(row.get("gate", {}).get("status") in ("pass", "not_applicable") for row in retained):
        status = "awaiting_ai_confirmation"
    elif result["summary"]["pending_count"]:
        status = "incomplete" if (run and run.get("stage") != "ai") or any(row.get("quality") for row in retained) else "ready"
    elif result["reward_ready"]:
        status = "confirmed"
    else:
        status = "awaiting_confirmation"
    result["status"] = status
    result["workflow"] = workflow_state(task)
    result["annotation"] = annotation_status(task)
    if (result.get("run") or {}).get("stage") == "quality":
        # Recheck cached progress for display only; never rewrite a saved run on read.
        complete_ids = {row["id"] for row in task["feedback"] if stage_complete(task, row, "quality")}
        result["run"]["completed_ids"] = [identifier for identifier in result["run"].get("completed_ids", [])
                                            if identifier in complete_ids]
    result["annotation_issues"] = readiness_issues(task)
    result["confirmation_issues"] = readiness_issues(task, settlement=True)
    if result.get("roster_snapshot"):
        result["roster_snapshot"].pop("entries", None)
    if summary:
        for key in ("feedback", "feedback_versions", "events", "uploads", "headers", "duplicate_groups", "input_issues", "confirmation_issues", "roster_snapshot", "runs"):
            result.pop(key, None)
        return result
    result.pop("feedback_versions", None)
    for row in result["feedback"]:
        row["outcome"] = effective_result(task, row)
        row["answers"] = [{"column": index, "question": task["headers"][index], "question_zh": (task.get("headers_zh") or task["headers"])[index], "answer": _cell(row, index)}
                          for index in (task.get("mapping") or {}).get("open_text_cols", [])]
    result["players"] = aggregate_players(task)
    return result


def list_tasks(login, *, feature="annotate") -> list[dict]:
    require_permission(login, feature)
    access = record_scope(login)
    result = [task_view(task, login, summary=True, access=access) for task in store.list_records("tasks") if not task.get("deleted_at") and access.can_read(task)]
    result.sort(key=lambda item: item["created_at"], reverse=True)
    queue = sorted((item for item in store.list_records("tasks") if (item.get("run") or {}).get("status") == "queued"),
                   key=lambda item: (item["run"].get("schedule_at", item["run"]["queued_at"]), item["run"]["id"]))
    positions = {item["id"]: index + 1 for index, item in enumerate(queue)}
    for item in result:
        item["queue_position"] = positions.get(item["id"])
    return result
