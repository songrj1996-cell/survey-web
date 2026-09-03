"""Create and serve immutable DOCX exports for approved Interview Report V2."""

from __future__ import annotations

from io import BytesIO
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
from typing import Any

from docx import Document

from app.core.interview_v2_export import (
    EXPORT_FORMAT,
    EXPORT_MEDIA_TYPE,
    EXPORT_PROFILE_VERSION,
    InterviewV2ExportValidationError,
    build_export_package,
    validate_export_visible_text,
)
from app.core.interview_v2_report import (
    InterviewV2ReportValidationError,
    validate_report_approval,
)
from app.core.security import _owner_from_login, _visible_to_owner
from app.services.interview_v2_import_service import InterviewV2ImportError
from app.services.interview_v2_report_review_service import _report_input
from app.services.interview_v2_report_service import (
    _is_report_current,
    _load_accessible_report,
)
from app.services.report_render import markdown_to_docx
from app.storage import interview_v2_store as store


_PUBLIC_ARTIFACT_FIELDS = (
    "export_artifact_id",
    "project_id",
    "report_version_id",
    "report_version_number",
    "format",
    "export_profile",
    "manifest",
    "manifest_sha256",
    "report_revision_payload_sha256",
    "content_sha256",
    "byte_size",
    "file_name",
    "media_type",
    "created_at",
)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def _payload_sha256(value: Any) -> str:
    encoded = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _validate_rendered_docx_privacy(content: bytes) -> None:
    """Scan the final visible DOCX text after Markdown formatting is removed."""

    document = Document(BytesIO(content))
    visible_parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            visible_parts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        for container in (section.header, section.footer):
            visible_parts.extend(paragraph.text for paragraph in container.paragraphs)
            for table in container.tables:
                for row in table.rows:
                    visible_parts.extend(cell.text for cell in row.cells)
    validate_export_visible_text("\n".join(visible_parts))


def _error(
    code: str,
    message: str,
    *,
    status: int,
    retryable: bool = False,
    suggested_action: str,
) -> InterviewV2ImportError:
    return InterviewV2ImportError(
        status_code=status,
        code=code,
        message=message,
        retryable=retryable,
        suggested_action=suggested_action,
    )


def _request_profile(request: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(request, dict) or not set(request) <= {
        "format",
        "include_evidence_appendix",
    }:
        raise _error(
            "EXPORT_REQUEST_INVALID",
            "导出请求格式无效。",
            status=400,
            suggested_action="fix_export_request",
        )
    export_format = request.get("format", "docx")
    include_appendix = request.get("include_evidence_appendix", True)
    if export_format != EXPORT_FORMAT or include_appendix is not True:
        raise _error(
            "EXPORT_REQUEST_INVALID",
            "当前仅支持包含证据附录的 DOCX 正式导出。",
            status=400,
            suggested_action="fix_export_request",
        )
    return {
        "format": EXPORT_FORMAT,
        "include_evidence_appendix": True,
        "export_profile_version": EXPORT_PROFILE_VERSION,
    }


def _public(artifact: dict[str, Any]) -> dict[str, Any]:
    result = {
        field: deepcopy(artifact.get(field))
        for field in _PUBLIC_ARTIFACT_FIELDS
    }
    result["status"] = "READY"
    result["download_url"] = (
        "/api/v1/interview-export-artifacts/"
        f"{artifact.get('export_artifact_id')}/download"
    )
    return result


def _load_evidence_by_id(
    project_id: str, report_revision: dict[str, Any]
) -> dict[str, dict[str, Any]]:
    analysis_source = (
        (report_revision.get("source") or {}).get("analysis_source") or {}
    )
    evidence_revision_id = str(
        analysis_source.get("evidence_revision_id") or ""
    )
    try:
        evidence_revision = store.load_evidence_revision(
            project_id, evidence_revision_id
        )
    except (OSError, TypeError, ValueError) as exc:
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "报告证据版本读取失败，请稍后重试。",
            status=500,
            retryable=True,
            suggested_action="retry_export",
        ) from exc
    if evidence_revision is None:
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "报告引用的证据版本不存在。",
            status=500,
            retryable=True,
            suggested_action="retry_export",
        )
    try:
        entries = store._evidence_entries(evidence_revision)
        evidence_by_id = {
            str(item.get("evidence_id") or ""): deepcopy(item)
            for item in entries
        }
    except (TypeError, ValueError) as exc:
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "报告证据版本完整性校验失败。",
            status=500,
            retryable=True,
            suggested_action="retry_export",
        ) from exc
    if len(evidence_by_id) != len(entries) or "" in evidence_by_id:
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "报告证据版本存在重复或无效标识。",
            status=500,
            retryable=True,
            suggested_action="retry_export",
        )
    return evidence_by_id


def _load_accessible_artifact(
    export_artifact_id: str, login: dict[str, Any] | None
) -> dict[str, Any]:
    try:
        store.validate_resource_id(export_artifact_id, "export")
    except (TypeError, ValueError) as exc:
        raise _error(
            "EXPORT_REQUEST_INVALID",
            "导出产物 ID 格式无效。",
            status=400,
            suggested_action="fix_export_request",
        ) from exc
    try:
        locator = store.locate_export_artifact(export_artifact_id)
    except ValueError as exc:
        # Locator corruption must not disclose artifact existence before owner
        # authentication. Treat it as absent for every caller.
        raise _error(
            "INTERVIEW_EXPORT_NOT_FOUND",
            "未找到该导出产物。",
            status=404,
            suggested_action="return_to_report_review",
        ) from exc
    except (OSError, TypeError) as exc:
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "导出产物定位失败，请稍后重试。",
            status=500,
            retryable=True,
            suggested_action="retry_export_download",
        ) from exc
    if locator is None:
        raise _error(
            "INTERVIEW_EXPORT_NOT_FOUND",
            "未找到该导出产物。",
            status=404,
            suggested_action="return_to_report_review",
        )
    try:
        project = store.load_project(str(locator.get("project_id") or ""))
    except (OSError, TypeError, ValueError) as exc:
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "导出产物所属项目读取失败。",
            status=500,
            retryable=True,
            suggested_action="retry_export_download",
        ) from exc
    if project is None or not _visible_to_owner(project, login):
        raise _error(
            "INTERVIEW_EXPORT_NOT_FOUND",
            "未找到该导出产物。",
            status=404,
            suggested_action="return_to_report_review",
        )
    try:
        saved = store.load_export_artifact(export_artifact_id)
    except (OSError, TypeError, ValueError) as exc:
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "导出产物完整性校验失败。",
            status=500,
            retryable=True,
            suggested_action="retry_export_download",
        ) from exc
    if (
        saved is None
        or saved.get("project_id") != locator.get("project_id")
    ):
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "导出产物定位与元数据不一致。",
            status=500,
            retryable=True,
            suggested_action="retry_export_download",
        )
    return saved


def create_export(
    report_version_id: str,
    request: dict[str, Any],
    login: dict[str, Any] | None,
) -> dict[str, Any]:
    """Render once and publish a current approved report as immutable DOCX."""

    request_profile = _request_profile(request)
    saved_report = _load_accessible_report(report_version_id, login)
    revision = saved_report["revision"]
    is_current_version = (
        (saved_report.get("state") or {}).get("current_report_version_id")
        == report_version_id
    )
    if revision.get("status") != "approved" or not is_current_version:
        raise _error(
            "REPORT_EXPORT_BLOCKED",
            "只有当前已批准的报告版本可以正式导出。",
            status=409,
            suggested_action="return_to_report_review",
        )
    try:
        upstream_current = _is_report_current(
            str(saved_report.get("project_id") or ""), revision
        )
    except (OSError, TypeError, ValueError) as exc:
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "报告上游状态读取失败，请稍后重试。",
            status=500,
            retryable=True,
            suggested_action="retry_export",
        ) from exc
    if not upstream_current:
        raise _error(
            "REPORT_EXPORT_BLOCKED",
            "报告引用的上游版本已变化，不能正式导出。",
            status=409,
            suggested_action="refresh_report",
        )
    try:
        validate_report_approval(
            revision, report_input=_report_input(revision)
        )
    except InterviewV2ReportValidationError as exc:
        raise _error(
            "REPORT_EXPORT_BLOCKED",
            "报告未通过导出前确定性复核，不能正式导出。",
            status=409,
            suggested_action="return_to_report_review",
        ) from exc

    request_fingerprint = _payload_sha256(
        {
            "report_version_id": report_version_id,
            "report_revision_payload_sha256": revision.get(
                "revision_payload_sha256"
            ),
            **request_profile,
        }
    )
    export_artifact_id = f"export_{request_fingerprint[:32]}"
    try:
        existing = store.load_export_artifact_bytes(export_artifact_id)
    except (OSError, TypeError, ValueError) as exc:
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "既有导出产物读取失败，请稍后重试。",
            status=500,
            retryable=True,
            suggested_action="retry_export",
        ) from exc
    if existing is not None:
        artifact = existing.get("artifact") or {}
        if (
            existing.get("project_id")
            != saved_report.get("project_id")
            or artifact.get("request_fingerprint") != request_fingerprint
            or artifact.get("report_revision_payload_sha256")
            != revision.get("revision_payload_sha256")
        ):
            raise _error(
                "EXPORT_PERSISTENCE_FAILED",
                "导出幂等标识发生冲突。",
                status=500,
                suggested_action="retry_export",
            )
        try:
            # Re-enter the store's project lock before returning an existing
            # artifact. This keeps idempotent retries subject to the same
            # current-report and upstream-source CAS as a first publication,
            # while the store reuses the exact persisted DOCX bytes.
            saved_existing = store.save_export_artifact(
                project_id=str(saved_report.get("project_id") or ""),
                artifact=artifact,
                content=existing["content"],
            )
        except store.ExportInputConflictError as exc:
            raise _error(
                "EXPORT_INPUT_CHANGED",
                "重复导出期间报告或上游版本已变化，请刷新后重试。",
                status=409,
                suggested_action="refresh_report",
            ) from exc
        except (FileExistsError, KeyError, OSError, TypeError, ValueError) as exc:
            raise _error(
                "EXPORT_PERSISTENCE_FAILED",
                "既有导出产物复核失败，请稍后重试。",
                status=500,
                retryable=True,
                suggested_action="retry_export",
            ) from exc
        return _public(saved_existing["artifact"])

    evidence_by_id = _load_evidence_by_id(
        str(saved_report.get("project_id") or ""), revision
    )
    try:
        package = build_export_package(
            revision,
            evidence_by_id,
            is_current_version=is_current_version,
        )
    except InterviewV2ExportValidationError as exc:
        raise _error(
            "REPORT_EXPORT_BLOCKED",
            "报告或证据未通过正式导出校验。",
            status=409,
            suggested_action="return_to_report_review",
        ) from exc
    try:
        markdown = package.get("markdown")
        if not isinstance(markdown, str) or not markdown:
            raise ValueError("export package markdown is invalid")
        content = markdown_to_docx(markdown)
        if not isinstance(content, bytes) or not content:
            raise ValueError("DOCX renderer returned invalid bytes")
        _validate_rendered_docx_privacy(content)
    except InterviewV2ExportValidationError as exc:
        raise _error(
            "REPORT_EXPORT_BLOCKED",
            "最终 Word 内容未通过导出隐私校验。",
            status=409,
            suggested_action="return_to_report_review",
        ) from exc
    except Exception as exc:
        raise _error(
            "EXPORT_RENDER_FAILED",
            "DOCX 导出渲染失败，请稍后重试。",
            status=500,
            retryable=True,
            suggested_action="retry_export",
        ) from exc

    now = _now()
    export_profile = deepcopy(package.get("export_profile"))
    manifest = deepcopy(package.get("manifest"))
    artifact = {
        "export_artifact_id": export_artifact_id,
        "project_id": saved_report.get("project_id"),
        "report_version_id": report_version_id,
        "report_version_number": revision.get("version_number"),
        "report_revision_payload_sha256": revision.get(
            "revision_payload_sha256"
        ),
        "request_fingerprint": request_fingerprint,
        "status": "READY",
        "format": EXPORT_FORMAT,
        "export_profile": export_profile,
        "export_profile_sha256": _payload_sha256(export_profile),
        "manifest": manifest,
        "manifest_sha256": package.get("manifest_sha256"),
        "file_name": package.get("filename"),
        "media_type": EXPORT_MEDIA_TYPE,
        "created_at": now,
        "ready_at": now,
        "created_by": str(
            _owner_from_login(login).get("owner_key") or ""
        ),
    }
    try:
        saved_artifact = store.save_export_artifact(
            project_id=str(saved_report.get("project_id") or ""),
            artifact=artifact,
            content=content,
        )
    except store.ExportInputConflictError as exc:
        raise _error(
            "EXPORT_INPUT_CHANGED",
            "导出生成期间报告或上游版本已变化，请刷新后重试。",
            status=409,
            suggested_action="refresh_report",
        ) from exc
    except FileExistsError as exc:
        try:
            concurrent = store.load_export_artifact_bytes(export_artifact_id)
        except (OSError, TypeError, ValueError):
            concurrent = None
        concurrent_artifact = (concurrent or {}).get("artifact") or {}
        if (
            concurrent_artifact.get("request_fingerprint")
            == request_fingerprint
            and concurrent_artifact.get("report_revision_payload_sha256")
            == revision.get("revision_payload_sha256")
        ):
            return _public(concurrent_artifact)
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "导出产物保存冲突，请稍后重试。",
            status=500,
            retryable=True,
            suggested_action="retry_export",
        ) from exc
    except (OSError, TypeError, ValueError) as exc:
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "导出产物保存失败，请稍后重试。",
            status=500,
            retryable=True,
            suggested_action="retry_export",
        ) from exc
    return _public(saved_artifact["artifact"])


def get_export_artifact(
    export_artifact_id: str, login: dict[str, Any] | None
) -> dict[str, Any]:
    saved = _load_accessible_artifact(export_artifact_id, login)
    return _public(saved["artifact"])


def get_export_download(
    export_artifact_id: str, login: dict[str, Any] | None
) -> dict[str, Any]:
    saved = _load_accessible_artifact(export_artifact_id, login)
    try:
        downloaded = store.load_export_artifact_bytes(export_artifact_id)
    except (OSError, TypeError, ValueError) as exc:
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "导出文件完整性校验失败，请稍后重试。",
            status=500,
            retryable=True,
            suggested_action="retry_export_download",
        ) from exc
    if (
        downloaded is None
        or downloaded.get("project_id") != saved.get("project_id")
        or downloaded.get("artifact") != saved.get("artifact")
    ):
        raise _error(
            "EXPORT_PERSISTENCE_FAILED",
            "导出文件与元数据不一致。",
            status=500,
            retryable=True,
            suggested_action="retry_export_download",
        )
    artifact = saved["artifact"]
    return {
        "artifact": _public(artifact),
        "content": downloaded["content"],
        "file_name": artifact["file_name"],
        "media_type": artifact["media_type"],
    }


__all__ = [
    "create_export",
    "get_export_artifact",
    "get_export_download",
]
