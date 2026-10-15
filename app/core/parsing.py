"""core/parsing:上传文件解析(CSV / Excel → 行数据)。纯无状态工具,不含业务逻辑。"""
import csv
import io
import zipfile

import openpyxl


def _parse_csv(content: bytes) -> list[list]:
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb2312"):
        try:
            text = content.decode(enc)
            reader = csv.reader(io.StringIO(text))
            rows = [list(row) for row in reader]
            while rows and all(not c.strip() for c in rows[-1]):
                rows.pop()
            return rows
        except (UnicodeDecodeError, csv.Error):
            continue
    raise ValueError("无法解析 CSV 文件，请确认文件编码为 UTF-8 或 GBK")


def _parse_excel(content: bytes) -> list[list]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            ws = wb.active
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append([("" if c is None else str(c)) for c in row])
        finally:
            wb.close()
    except (
        OSError,
        ValueError,
        KeyError,
        zipfile.BadZipFile,
        openpyxl.utils.exceptions.InvalidFileException,
    ) as exc:
        raise ValueError("无法解析 .xlsx 文件，请确认文件未损坏且格式正确") from exc
    while rows and all(not c.strip() for c in rows[-1]):
        rows.pop()
    return rows


def _parse_file(filename: str, content: bytes) -> list[list]:
    name_lower = filename.lower()
    if name_lower.endswith(".csv"):
        return _parse_csv(content)
    elif name_lower.endswith(".xlsx"):
        return _parse_excel(content)
    raise ValueError("不支持的文件格式，请上传 .csv 或 .xlsx 文件")


def parse_file_with_cell_metadata(filename: str, content: bytes) -> tuple[list[list], dict]:
    """Opt-in identity safety metadata; existing survey parsing is unchanged."""
    rows = _parse_file(filename, content)
    issues: dict[str, str] = {}
    if filename.lower().endswith(".xlsx"):
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            for row_index, cells in enumerate(workbook.active.iter_rows()):
                for column_index, cell in enumerate(cells):
                    value = cell.value
                    issue = ""
                    if cell.data_type == "f":
                        issue = "身份单元格是公式，请改用原始文本 ID"
                    elif isinstance(value, (int, float)) and not isinstance(value, bool):
                        if isinstance(value, float) and (not value.is_integer() or abs(value) >= 10 ** 15):
                            issue = "数字身份可能已失去精度，请从原始来源补填文本 ID"
                        elif abs(value) >= 10 ** 15:
                            issue = "长 ID 保存为 Excel 数字，无法证明精度完整，请补填文本 ID"
                    if issue:
                        issues[f"{row_index}:{column_index}"] = issue
        finally:
            workbook.close()
    return rows, issues
