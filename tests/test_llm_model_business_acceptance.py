"""Published settings -> business entry -> fake HTTP -> result/usage/log checks.

Only the gateway is simulated. No app.main, real credentials or external HTTP.
V2 project-read fixtures stand in for its separately tested import workflow.
"""
import asyncio
from contextlib import ExitStack, redirect_stdout
from copy import deepcopy
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

import httpx
from fastapi import FastAPI
from starlette.requests import Request

from app.core import config, security
from app.core.llm_context import bind_llm_api_key, bind_llm_attempt_observer
from app.core.llm_model_registry import environment_defaults
from app.integrations import llm_client
from app.routers import settings_api
from app.services import auth, llm_credentials, llm_model_settings as settings, llm_usage
from app.storage import history as history_store, llm_usage as usage_store, sessions
from app.storage import glossary as glossary_store, prompts as prompt_store


LOGIN = {"email": "acceptance@example.test", "name": "离线验收员"}
FIXTURE_KEY = "offline-acceptance-fixture"
PRIVATE_MATERIAL = "PRIVATE_ACCEPTANCE_MATERIAL_完整材料末尾"


def configuration(generation="a"):
    result = environment_defaults()
    result["models"] = []
    for stage, task in result["tasks"].items():
        candidates = []
        for role, effort, cap in (("primary", "high", 4096), ("backup", "low", 2048)):
            model = "gpt-accept-" + stage.replace(".", "-") + "-" + generation + "-" + role
            result["models"].append({"id": model, "name": model,
                "efforts": ["default", "low", "high"], "context_tokens": 100000,
                "max_output_tokens": 64000, "verified": True,
                "verification_note": "Synthetic offline fixture, not a real capability claim"})
            candidates.append({"model": model, "reasoning": effort, "max_tokens": cap})
        task.update(primary=candidates[0], fallbacks=candidates[1:],
                    timeout_seconds=5, max_http_attempts=3)
    return result


def decode_events(chunks):
    return [json.loads(chunk[6:].strip()) for chunk in chunks if chunk.startswith("data: ")]


class AcceptanceCase(unittest.IsolatedAsyncioTestCase):
    """Shared fixture, deliberately no test methods (also used by restart tests)."""

    async def asyncSetUp(self):
        asyncio.get_running_loop().set_debug(False)
        root = Path(__file__).resolve().parents[1] / ".test-artifacts" / "model-management"
        self.assertTrue(root.is_dir(), "Use the already approved isolated artifact directory")
        # V2 resource IDs make paths long on Windows; use the native extended path
        # for both writes AND TemporaryDirectory cleanup. This is fixture setup,
        # not evidence that ordinary production paths support the same lengths.
        if os.name == "nt":
            root = Path("\\\\?\\" + str(root))
        self.directory = tempfile.TemporaryDirectory(prefix="acceptance-", dir=root)
        self.assertTrue(Path(self.directory.name).resolve().is_relative_to(root.resolve()))
        self.addCleanup(self.directory.cleanup)
        self.folder = Path(self.directory.name)
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.logs = io.StringIO()
        self.stack.enter_context(redirect_stdout(self.logs))
        for target, key, value in (
            (config, "DATA_DIR", str(self.folder)),
            (config, "LLM_MODEL_SETTINGS_FILE", str(self.folder / "models.json")),
            (config, "INTERVIEW_V2_DATA_DIR", str(self.folder / "iv2")),
            (sessions, "_SESSION_DIR", self.folder / "sessions"),
            (history_store, "HISTORY_FILE", str(self.folder / "history.json")),
            (usage_store, "USER_LLM_USAGE_FILE", str(self.folder / "usage.json")),
            (prompt_store, "PROMPTS_FILE", str(self.folder / "prompts.json")),
            (glossary_store, "GLOSSARY_FILE", str(self.folder / "glossary.json")),
            (security, "FEISHU_ADMIN_EMAILS", {LOGIN["email"]}),
            (security, "FEISHU_ALLOWED_EMAILS", set()),
            (llm_client, "LLM_API_BASE", "https://model-acceptance.invalid/v1"),
            (llm_client, "LLM_REPORT_MAX_ATTEMPTS", 1),
        ):
            self.stack.enter_context(patch.object(target, key, value))
        self.stack.enter_context(patch.object(auth, "_current_login", AsyncMock(return_value=LOGIN)))
        self.stack.enter_context(patch.object(llm_credentials, "_current_login", AsyncMock(return_value=LOGIN)))
        self.usage_creation_errors = []
        start_usage = llm_credentials.start_llm_usage_task
        def capture_usage_start(*args, **kwargs):
            try:
                return start_usage(*args, **kwargs)
            except Exception as exc:
                self.usage_creation_errors.append(type(exc).__name__ + ": " + str(exc))
                raise
        self.stack.enter_context(patch.object(llm_credentials, "start_llm_usage_task", capture_usage_start))
        self.stack.enter_context(patch.object(settings_api, "audit_log", AsyncMock()))
        self.stack.enter_context(patch.object(tempfile, "tempdir", str(self.folder)))
        self.calls = []
        self.answers = {}
        self.failures = {}
        self.started = asyncio.Event()
        self.release = asyncio.Event()
        self.block_stage = None

        async def gateway(transport, request):
            self.assertEqual(request.url.host, "model-acceptance.invalid", "External HTTP is forbidden")
            self.assertTrue(str(request.url).endswith("/responses"))
            payload = json.loads(request.content)
            self.calls.append(deepcopy(payload))
            model = payload["model"]
            stage = self.stage_for_model(model)
            if self.block_stage == stage:
                self.started.set()
                await self.release.wait()
            mode = self.failures.get(stage)
            if mode == "all503" or (mode in {"503", "timeout"} and model.endswith("primary")):
                if mode == "timeout":
                    raise httpx.ReadTimeout("synthetic timeout", request=request)
                return httpx.Response(503, json={"error": {"message": "synthetic unavailable"}})
            answer = self.answers[stage]
            if callable(answer):
                answer = answer(payload)
            if not isinstance(answer, str):
                answer = json.dumps(answer, ensure_ascii=False)
            completed = {"status": "completed", "model": model}
            if mode != "unknown":
                completed["usage"] = {"input_tokens": 11, "output_tokens": 7, "total_tokens": 18}
            events = [{"type": "response.output_text.delta", "delta": answer},
                      {"type": "response.completed", "response": completed}]
            body = "".join("data: " + json.dumps(event, ensure_ascii=False) + "\n\n" for event in events)
            return httpx.Response(200, headers={"content-type": "text/event-stream"}, content=body.encode())

        self.stack.enter_context(patch.object(httpx.AsyncHTTPTransport, "handle_async_request", gateway))
        app = FastAPI()
        app.include_router(settings_api.router)
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://isolated.test")
        self.addAsyncCleanup(self.client.aclose)
        self.state = None
        await self.publish("a")

    async def publish(self, generation):
        state = (await self.client.get("/api/model-settings")).json()
        response = await self.client.post("/api/model-settings/publish", json={
            "expected_revision": state["revision"], "configuration": configuration(generation)})
        self.assertEqual(response.status_code, 200, response.text)
        self.state = response.json()
        return await settings.build_model_snapshot()

    def stage_for_model(self, model):
        for stage in configuration()["tasks"]:
            if model.startswith("gpt-accept-" + stage.replace(".", "-") + "-"):
                return stage
        self.fail("Unexpected model: " + model)

    def session(self, **fields):
        sid = sessions.new_session()
        sess = {"session_id": sid, **fields}
        sessions.save_session(sid, sess)
        return sid

    async def consume(self, stream, *, sid="", category="survey"):
        request = Request({"type": "http", "method": "GET", "path": "/offline-acceptance", "headers": []})
        wrapped = llm_credentials.stream_with_llm_api_key(stream, FIXTURE_KEY, request=request,
            category=category, action="离线业务验收", reference_id=sid)
        return decode_events([chunk async for chunk in wrapped])

    def assert_usage(self, expected_stages, *, missing=0, version=1):
        self.assertEqual(self.usage_creation_errors, [], "Usage task creation failed")
        page = llm_usage.get_user_llm_usage(LOGIN, period="all")
        record = page["records"][0]
        used = record["model_settings_used"]
        self.assertEqual({item["model_task_key"] for item in used}, set(expected_stages))
        self.assertEqual({item["model_config_version"] for item in used}, {version})
        self.assertEqual(record["call_count"], len(self.calls))
        self.assertEqual(record["usage_missing_call_count"], missing)
        self.assertEqual(sum(item["http_attempts"] for item in used), len(self.calls))
        self.assertEqual({item["model"] for item in used}, {item["model"] for item in self.calls})
        for item in used:
            backup = item["model"].endswith("backup")
            self.assertEqual(item["reasoning_effort"], "low" if backup else "high")
            self.assertEqual(item["max_output_tokens"], 2048 if backup else 4096)
            self.assertEqual(item["protocol"], "responses")
        for request in self.calls:
            backup = request["model"].endswith("backup")
            self.assertEqual(request["reasoning"]["effort"], "low" if backup else "high")
            self.assertEqual(request["max_output_tokens"], 2048 if backup else 4096)
        logs = self.logs.getvalue()
        for secret in (FIXTURE_KEY, PRIVATE_MATERIAL, LOGIN["email"]):
            self.assertNotIn(secret, logs)
        self.assertIn("配置版本", logs)
        for model in {item["model"] for item in self.calls}:
            self.assertIn(model, logs)
        self.assertRegex(logs, r"\d{4}-\d\d-\d\d .*\+08:00")
        return record

    def annotation_quality_case(self):
        from app.services import annotate_workflow as service
        sid = "offline-quality-acceptance"
        sess = {"filename": "synthetic-quality.xlsx", "headers": ["ID", "请描述排队体验"],
                "rows": [["ID", "请描述排队体验"], ["P1", "The queue takes ten minutes."],
                         ["P2", "I cannot find a match at night."]],
                "id_col": 0, "open_text_cols": [1], "tasks": {"ai_detect": False, "quality": True},
                "background": PRIVATE_MATERIAL}
        self.stack.enter_context(patch.object(service, "annotate_sessions", {sid: sess}))
        self.stack.enter_context(patch.object(service, "ANNOTATE_RESULT_DIR", self.folder / "annotated"))
        self.stack.enter_context(patch.object(service, "_current_login", AsyncMock(return_value=None)))
        self.stack.enter_context(patch.object(service, "audit_log", AsyncMock()))
        results = [{"id": row[0], "q_labels": {"col_1": "有效反馈"},
            "q_reasons": {"col_1": "回答描述了本人匹配等待的具体体验，满足题目要求"},
            "q_evidence": {"col_1": row[1]},
            "q_checks": {"col_1": {"requirement": "specific_description", "support": "substantive"}},
            "overall": "有效反馈", "overall_reason": "提供了本人匹配体验的有效信息", "translations": {}}
            for row in sess["rows"][1:]]
        self.answers["annotate.quality"] = deepcopy(results)
        self.answers["annotate.translation"] = [
            {"id": "P1", "key": "col_1", "translation": "排队需要十分钟。"},
            {"id": "P2", "key": "col_1", "translation": "我在晚上匹配不到对局。"}]
        return service, sid, sess, results

    def quick_report_case(self, *, two_questions=False):
        from app.services import survey_service as service
        from tests.test_report_quick_mode_flow import _quick_session
        from tests.test_report_quick_pipeline import answer
        fixture = _quick_session(two_questions=two_questions)
        fixture["owner_key"] = "email:" + LOGIN["email"]
        sid = self.session(**fixture)
        self.stack.enter_context(patch.object(service, "_current_login", AsyncMock(return_value=LOGIN)))
        self.stack.enter_context(patch.object(service, "is_quick_report_enabled", return_value=True))
        self.stack.enter_context(patch.object(service, "audit_log", AsyncMock()))
        self.answers["quick.summary"] = lambda request: answer(json.loads(request["input"][-1]["content"]))
        return service, sid


class BusinessAcceptanceTests(AcceptanceCase):
    async def test_annotation_quality_translation_fallback_and_workbook_archive(self):
        from openpyxl import load_workbook
        service, sid, sess, _ = self.annotation_quality_case()
        self.failures.update({"annotate.quality": "503", "annotate.translation": "503"})
        events = await self.consume(service.quality_stream(sid, object()), sid=sid, category="annotate")
        done = next(event for event in events if event["type"] == "quality_done")
        self.assertTrue(done["history_saved"])
        self.assertEqual(done["missing_ids"], [])
        self.assertEqual(done["missing_translation_ids"], [])
        self.assertEqual(sess["quality_status"], "complete")
        self.assertEqual(len(self.calls), 4)
        for stage in ("annotate.quality", "annotate.translation"):
            requests = [c for c in self.calls if self.stage_for_model(c["model"]) == stage]
            self.assertEqual(requests[0]["input"], requests[1]["input"])
            self.assertTrue(requests[0]["model"].endswith("primary"))
            self.assertTrue(requests[1]["model"].endswith("backup"))
        self.assert_usage(["annotate.quality", "annotate.translation"], missing=2)
        archived = history_store._load_history()[0]
        self.assertEqual(archived["llm_model_stage_snapshots"]["quality"]["llm_model_snapshot"]["version"], 1)
        path = Path(archived["annotate_result_path"])
        self.assertTrue(path.is_relative_to(self.folder))
        workbook = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
        try:
            cells = [str(value) for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row]
        finally:
            workbook.close()
        for expected in ("The queue takes ten minutes.", "排队需要十分钟。", "有效反馈"):
            self.assertIn(expected, cells)

    async def test_quick_report_service_archives_snapshot_sources_and_usage(self):
        service, sid = self.quick_report_case()
        events = await self.consume(service.report_stream(sid, object()), sid=sid)
        self.assertIn("report_done", [event["type"] for event in events])
        saved = sessions.get_session(sid)
        version = saved["report_versions"][-1]
        self.assertEqual(version["report_status"], "complete")
        self.assertEqual(version["llm_model_snapshot"]["version"], 1)
        self.assertEqual(version["quick_checkpoint"]["llm_model_snapshot"]["version"], 1)
        self.assertEqual(len(self.calls), 1)
        archived = history_store._load_history()[0]
        self.assertEqual(archived["report_versions"], saved["report_versions"])
        self.assertEqual(archived["input_snapshot"], version["input_snapshot"])
        self.assert_usage(["quick.summary"])

    async def test_all_38_published_stages_reach_transport_with_independent_options(self):
        self.answers = {key: "synthetic answer" for key in configuration()["tasks"]}
        recorder = llm_usage.start_llm_usage_task(LOGIN, category="other", action="38 环节验收")
        async with settings.model_operation_scope():
            with bind_llm_api_key(FIXTURE_KEY), bind_llm_attempt_observer(recorder.on_attempt_event):
                for stage in self.answers:
                    answer, model = await settings.call_model(llm_client.collect_chat_completion, stage,
                        [{"role": "user", "content": PRIVATE_MATERIAL}])
                    self.assertEqual(answer, "synthetic answer")
                    self.assertEqual(self.stage_for_model(model), stage)
        recorder.finish("completed")
        self.assertEqual(len(self.calls), 38)
        self.assert_usage(self.answers)
        self.assertTrue(all(PRIVATE_MATERIAL in json.dumps(c, ensure_ascii=False) for c in self.calls))

    async def test_survey_and_quant_plans_use_their_published_stage(self):
        from app.services import survey_service as service
        from tests.test_survey_planner_direct_llm import _plan_json
        for mode, stage in (("standard", "survey.plan"), ("crosstab", "quant.plan")):
            with self.subTest(mode=mode):
                self.calls.clear()
                self.answers[stage] = (_plan_json() if mode == "standard" else {
                    "parts": [{"name": "总体评价", "scope": "满意度"}], "open_questions": []})
                sid = self.session(mode=mode, rows=[["段位", "主玩位置", "界面满意度"], ["金", "中", "4"]],
                    confirmed_columns=[], questionnaire_text=PRIVATE_MATERIAL, crosstab_questions=["满意度"])
                with patch.object(service, "audit_log", AsyncMock()), \
                     patch.object(service, "_ensure_branch_rules", return_value=[]):
                    events = await self.consume(service.plan_stream(sid, object()), sid=sid)
                self.assertIn("plan_ready", [item["type"] for item in events])
                saved = sessions.get_session(sid)
                self.assertEqual(self.stage_for_model(saved["planner_model"]), stage)
                self.assert_usage([stage])

    async def test_completed_plan_retains_configuration_provenance(self):
        """MM-O02: planning keeps its own frozen configuration."""
        from app.services import survey_service as service
        from tests.test_survey_planner_direct_llm import _plan_json
        self.answers["survey.plan"] = _plan_json()
        sid = self.session(rows=[["段位", "主玩位置", "界面满意度"], ["金", "中", "4"]])
        with patch.object(service, "audit_log", AsyncMock()), \
             patch.object(service, "_ensure_branch_rules", return_value=[]):
            events = await self.consume(service.plan_stream(sid, object()), sid=sid)
        self.assertIn("plan_ready", [item["type"] for item in events])
        self.assert_usage(["survey.plan"])
        # A plan's provenance must have its own slot; overwriting a report's
        # root snapshot would incorrectly change later report recovery.
        saved = sessions.get_session(sid)
        metadata = (saved.get("llm_model_stage_snapshots") or {}).get("plan") or {}
        self.assertEqual((metadata.get("llm_model_snapshot") or {}).get("version"), 1)

    async def test_initial_and_revised_plans_keep_stage_snapshot_without_replacing_report(self):
        from app.services import survey_service as service
        from tests.test_survey_planner_direct_llm import _plan_json
        report_settings = await settings.build_model_snapshot()
        await self.publish("b")
        for mode, stage in (("standard", "survey.plan"), ("quant", "quant.plan"), ("crosstab", "quant.plan")):
            with self.subTest(mode=mode):
                self.answers[stage] = (_plan_json() if mode != "crosstab" else {
                    "parts": [{"name": "总体评价", "scope": "满意度"}], "open_questions": []})
                report = {"version": 1, "report_md": "# 原有报告\n\n原正文与引用保持不变。",
                          "llm_model_snapshot": deepcopy(report_settings)}
                sid = self.session(mode=mode, rows=[["段位", "主玩位置", "界面满意度"], ["金", "中", "4"]],
                    confirmed_columns=[], questionnaire_text=PRIVATE_MATERIAL, crosstab_questions=["满意度"],
                    report_md=report["report_md"], report_versions=[report], active_report_version=1,
                    llm_model_snapshot=deepcopy(report_settings),
                    llm_model_stage_snapshots={"other": {"untouched": True}})
                for revision in (False, True):
                    self.calls.clear()
                    with patch.object(service, "audit_log", AsyncMock()), \
                         patch.object(service, "_ensure_branch_rules", return_value=[]):
                        stream = (service.plan_revision_stream(sid, "合成修订意见", object()) if revision
                                  else service.plan_stream(sid, object()))
                        events = await self.consume(stream, sid=sid)
                    self.assertIn("plan_ready", [event["type"] for event in events])
                    saved = sessions.get_session(sid)
                    self.assertEqual(saved["llm_model_stage_snapshots"]["plan"]["llm_model_snapshot"]["version"], 2)
                    self.assertEqual(saved["llm_model_snapshot"], report_settings)
                    self.assertEqual(saved["report_versions"], [report])
                    self.assertEqual(saved["report_md"], report["report_md"])
                    self.assertEqual(saved["llm_model_stage_snapshots"]["other"], {"untouched": True})
                    self.assertEqual(len(self.calls), 1, "Snapshot persistence must add no model requests")
                    self.assert_usage([stage], version=2)
                success_lines = [line for line in self.logs.getvalue().splitlines() if "本次规划配置已随方案保存" in line]
                self.assertTrue(success_lines)
                self.assertTrue(all("配置版本：v2" in line for line in success_lines))

    async def test_plan_snapshot_uses_operation_version_when_publication_happens_inflight(self):
        from app.services import survey_service as service
        from tests.test_survey_planner_direct_llm import _plan_json
        self.answers["survey.plan"] = _plan_json()
        self.block_stage = "survey.plan"
        sid = self.session(rows=[["段位", "主玩位置", "界面满意度"], ["金", "中", "4"]])
        with patch.object(service, "audit_log", AsyncMock()), \
             patch.object(service, "_ensure_branch_rules", return_value=[]):
            task = asyncio.create_task(self.consume(service.plan_stream(sid, object()), sid=sid))
            try:
                await asyncio.wait_for(self.started.wait(), timeout=10)
                await self.publish("b")
                self.block_stage = None
                self.release.set()
                events = await asyncio.wait_for(task, timeout=10)
            finally:
                self.release.set()
                task.cancel()
                await asyncio.gather(task, return_exceptions=True)
        self.assertIn("plan_ready", [event["type"] for event in events])
        saved = sessions.get_session(sid)
        self.assertEqual(saved["llm_model_stage_snapshots"]["plan"]["llm_model_snapshot"]["version"], 1)
        self.assertEqual((await settings.build_model_snapshot())["version"], 2)
        self.assert_usage(["survey.plan"])

    async def test_failed_planning_never_replaces_last_successful_stage_snapshot(self):
        from app.services import survey_service as service
        from tests.test_survey_planner_direct_llm import _plan_json
        self.answers["survey.plan"] = _plan_json()
        sid = self.session(rows=[["段位", "主玩位置", "界面满意度"], ["金", "中", "4"]])
        with patch.object(service, "audit_log", AsyncMock()), \
             patch.object(service, "_ensure_branch_rules", return_value=[]):
            await self.consume(service.plan_stream(sid, object()), sid=sid)
            original = sessions.get_session(sid)
            await self.publish("b")
            for cause in ("invalid_json", "upstream_failure", "save_failure"):
                with self.subTest(cause=cause):
                    self.answers["survey.plan"] = "not json" if cause == "invalid_json" else _plan_json()
                    self.failures.clear()
                    if cause == "upstream_failure":
                        self.failures["survey.plan"] = "all503"
                    start = len(self.logs.getvalue())
                    with ExitStack() as failure_scope:
                        if cause == "save_failure":
                            failure_scope.enter_context(patch.object(service, "save_session", side_effect=OSError("synthetic save failure")))
                        events = await self.consume(service.plan_revision_stream(sid, "合成修订意见", object()), sid=sid)
                    self.assertIn("error", [event["type"] for event in events])
                    self.assertNotIn("plan_ready", [event["type"] for event in events])
                    saved = sessions.get_session(sid)
                    self.assertEqual(saved["plan"], original["plan"])
                    self.assertEqual(saved["llm_model_stage_snapshots"], original["llm_model_stage_snapshots"])
                    self.assertNotIn("本次规划配置已随方案保存", self.logs.getvalue()[start:])

    async def test_comment_full_pipeline_and_quotes_archive_published_snapshot(self):
        from app.services import comment_pipeline as pipeline, comment_run
        prompts = {key: key for key in ("relevance", "extract", "merge", "classify", "report", "quote_batch", "quote_final")}
        for key in prompts:
            self.stack.enter_context(patch.object(pipeline, "_get_comment_" + key + "_system_prompt", return_value=key))
        self.answers.update({
            "comment.relevance": [{"idx": 0, "is_related": True, "relation": "direct", "reason": "相关"}],
            "comment.extract": [{"theme_name": "等待", "description": "等待体验", "sentiment": "negative"}],
            "comment.merge": [{"theme_id": "wait", "theme_name": "等待", "description": "等待体验", "sentiment": "negative"}],
            "comment.classify": [{"idx": 0, "theme_ids": ["wait"], "sentiment": "negative", "is_quote_candidate": True, "translation": "等待很久"}],
            "comment.report": "## 核心结论\n\n等待影响体验\n\n## 玩家核心观点\n\n等待很久\n\n## 业务建议\n\n核对等待流程",
        })
        def quotes(payload):
            query = json.loads(payload["input"][-1]["content"])
            return [{"idx": item["idx"], "text": item["text"], "translation": "等待很久",
                     "score": 90, "reason": "完整原文"} for item in query["comments_json"]]
        self.answers.update({"comment.quote_batch": quotes, "comment.quote_final": quotes})
        sid = self.session(kind="comment", mode="comment", comment_post_title="等待体验",
            comment_post_content="等待", comment_sample=["Waiting too long " + PRIVATE_MATERIAL],
            comment_long_candidates=["Waiting too long " + PRIVATE_MATERIAL], comment_sample_meta={})
        events = await self.consume(comment_run.comment_run_stream(sid, object()), sid=sid, category="comment")
        self.assertIn("comment_done", [item["type"] for item in events])
        self.assertIn("comment_quotes_done", [item["type"] for item in events])
        saved = sessions.get_session(sid)
        self.assertEqual(len(saved["comment_selected_raw_comments"]), 1)
        self.assertEqual(saved["llm_model_snapshot"]["version"], 1)
        history = history_store._load_history()
        self.assertEqual(history[0]["llm_model_snapshot"]["version"], 1)
        self.assert_usage(["comment." + key for key in prompts])

    async def test_interview_extract_write_audit_and_archive(self):
        from app.core.interview_parsing import parse_interview_workbook, serialize_interview_workbook
        from app.services import interview_service as service
        from tests.test_interview_report import _workbook_bytes, _Request
        workbook = parse_interview_workbook("访谈.xlsx", _workbook_bytes())
        self.answers.update({
            "interview.extract": {"players": [{"player_id": "P01", "aliases": ["玩家甲"]}],
                "modules": [{"title": "战斗反馈", "evidence": [{"player_id": "P01",
                    "need": "理解失败原因", "finding": "失败反馈不清楚", "record_excerpt": "希望失败后能立刻知道原因",
                    "source_refs": ["记录者A!B2"]}]}], "limitations": []},
            "interview.report": "## 战斗反馈\n\n### 模块判断\n\n需要理解失败原因。\n\n### 主要发现\n\n#### 发现1：反馈不清楚\n\n- P01：希望失败后能立刻知道原因。[来源：记录者A!B2]\n\n### 产品建议\n\n增加诊断提示。",
            "interview.audit": {"ok": True, "issues": []},
        })
        sid = self.session(kind="interview", mode="interview", filename="访谈.xlsx",
            interview_workbook=workbook, interview_source_text=serialize_interview_workbook(workbook))
        with patch.object(service, "_visible_to_owner", return_value=True):
            events = await self.consume(service.interview_report_stream(sid, _Request()), sid=sid, category="interview")
        self.assertIn("interview_done", [item["type"] for item in events])
        saved = sessions.get_session(sid)
        self.assertEqual(saved["interview_status"], "completed")
        self.assertEqual(saved["llm_model_snapshot"]["version"], 1)
        self.assertEqual(history_store._load_history()[0]["llm_model_snapshot"]["version"], 1)
        self.assert_usage(["interview.extract", "interview.report", "interview.audit"])

    async def test_announcement_entry_and_its_own_usage_recorder(self):
        from app.services import task_announcement as service
        self.answers["other.announcement"] = {"name": {"value": "活动甲", "evidence": "活动甲"},
            "base_reward": {"value": 100, "evidence": "基础奖励100钻"}}
        with patch.object(service, "require_background_llm_api_key", return_value=FIXTURE_KEY), \
             patch.object(service, "require_permission"):
            result = await service.extract_announcement("活动甲，基础奖励100钻。" + PRIVATE_MATERIAL, LOGIN)
        self.assertEqual(result["suggestions"]["base_reward"]["value"], 100)
        self.assert_usage(["other.announcement"])

    async def test_annotation_entry_preserves_stage_settings_in_archived_workbook(self):
        from app.services import annotate_workflow as service
        sid = "offline-model-acceptance"
        sess = {"filename": "synthetic.xlsx", "rows": [["ID", "反馈"], ["P1", "等待很久"]],
                "headers": ["ID", "反馈"], "id_col": 0, "open_text_cols": [1],
                "tasks": {"ai_detect": True, "quality": False}, "background": PRIVATE_MATERIAL}
        self.stack.enter_context(patch.object(service, "annotate_sessions", {sid: sess}))
        self.stack.enter_context(patch.object(service, "ANNOTATE_RESULT_DIR", self.folder / "annotated"))
        self.stack.enter_context(patch.object(service, "_current_login", AsyncMock(return_value=None)))
        self.stack.enter_context(patch.object(service, "audit_log", AsyncMock()))
        self.answers["annotate.ai"] = [{"id": "P1", "ai_prob": 10, "polish_prob": 5,
            "reason": "合成判断", "evidence": "", "counter_evidence": "等待很久",
            "translations": {"col_1": "等待很久"}}]
        events = await self.consume(service.ai_detect_stream(sid, object()), category="annotate")
        self.assertIn("ai_detect_done", [item["type"] for item in events])
        self.assertEqual(sess["ai_status"], "complete")
        self.assertEqual(sess["llm_model_stage_snapshots"]["ai"]["llm_model_snapshot"]["version"], 1)
        archived = history_store._load_history()[0]
        self.assertEqual(archived["llm_model_stage_snapshots"]["ai"]["llm_model_snapshot"]["version"], 1)
        self.assertTrue(Path(archived["annotate_result_path"]).is_relative_to(self.folder))
        self.assertTrue(Path(archived["annotate_result_path"]).is_file())
        self.assert_usage(["annotate.ai"])

    async def test_v2_dossier_entry_saves_both_stage_models_and_snapshot(self):
        from app.services import interview_v2_dossier_service as service
        from tests import test_interview_v2_dossier_service as fixture
        service.store._mapping_lock_path(fixture.PROJECT).parent.mkdir(parents=True)
        self.answers.update({
            "interview_v2.attribute": {"participant_id": fixture.PARTICIPANT,
                "facts": [{"candidate_id": "f1", "attribute_key": "frequency", "raw_value": "每天",
                    "fact_source": "explicit_self_report", "evidence_ids": [fixture.EV_BACKGROUND]}], "analytical_labels": []},
            "interview_v2.dossier": {"participant_id": fixture.PARTICIPANT,
                "claims": [{"claim_type": "behavior", "statement": "会使用该功能",
                    "supporting_evidence_ids": [fixture.EV_BODY], "conflicting_evidence_ids": []}],
                "contradictions": [], "missing_context": []},
        })
        recorder = llm_usage.start_llm_usage_task(LOGIN, category="interview", action="V2 档案验收")
        with patch.object(service, "_ready_project", return_value=fixture.ready_payload()), \
             bind_llm_api_key(FIXTURE_KEY), bind_llm_attempt_observer(recorder.on_attempt_event):
            result = await service.regenerate_dossier(fixture.PROJECT, fixture.PARTICIPANT,
                {"base_dossier_version_id": None}, LOGIN)
        recorder.finish("completed")
        self.assertEqual(result["status"], "generated")
        saved = service.store.load_current_participant_dossier(fixture.PROJECT, fixture.PARTICIPANT)
        self.assertEqual(saved["revision"]["llm_model_snapshot"]["version"], 1)
        self.assert_usage(["interview_v2.attribute", "interview_v2.dossier"])

    async def test_v2_continuous_project_import_through_report_approval(self):
        """Real synthetic workbook and storage, with no ready-project shortcut."""
        from openpyxl import Workbook
        from app.services import interview_v2_import_service as imports
        from app.services import interview_v2_mapping_service as mapping
        from app.services import interview_v2_structure_service as structure
        from app.services import interview_v2_analysis_boundary_service as boundary
        from app.storage import interview_v2_store as store
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "合成记录"
        sheet.append(["功能模块", "行类型", "问题/备注", "P01", "P02"])
        sheet.append(["组队", "模块", "", "", ""])
        sheet.append(["", "主问题", "1. 你如何组队？", "自己邀请好友", "随机匹配"])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        attempt, created = imports.create_upload_attempt(filename="synthetic-v2.xlsx", content=output.getvalue(),
            login=LOGIN, research_focus="组队方式", file_contract_version=config.INTERVIEW_V2_FILE_CONTRACT_VERSION,
            contract_acknowledged=True, idempotency_key="offline-continuous-import")
        self.assertTrue(created)
        checked = imports.run_upload_precheck(attempt["upload_attempt_id"])
        self.assertEqual(checked["status"], "ACCEPTED", checked)
        imported = imports.get_interview_import(checked["import_id"], LOGIN)
        import_id, project_id = imported["import_id"], imported["project_id"]
        proposals = mapping.get_group_proposals(import_id, LOGIN)
        saved = mapping.save_group_mapping(import_id, {"base_mapping_revision": 0, "groups": [{
            "display_name": "合成组", "sheets": [{"sheet_id": "sheet_001", "role": "record", "recorder_label": "合成记录员"}],
            "participants": [{"participant_label": label, "columns": [{"sheet_id": "sheet_001", "column_index": column}]}
                             for label, column in (("P01", 4), ("P02", 5))]}], "ignored_sheet_ids": []}, LOGIN)
        self.assertTrue(saved["confirmation_ready"], saved)
        confirmed = mapping.confirm_group_mapping(import_id, {"base_mapping_revision": saved["revision_number"],
            "mapping_sha256": saved["mapping_sha256"]}, LOGIN)
        built = structure.build_structure(import_id, {
            "base_mapping_revision_id": confirmed["mapping_revision_id"],
            "base_mapping_sha256": confirmed["mapping_sha256"]}, LOGIN)
        self.assertEqual(built["status"], "READY_FOR_DOSSIERS", built)
        proposed_boundary = boundary.get_analysis_boundary(import_id, LOGIN)
        self.assertTrue(proposed_boundary, "Boundary proposal must be available after actual import")
        chosen = deepcopy(proposed_boundary["analysis_boundary"])
        for field in ("evaluation_objects", "source_scope_rules", "label_scope_rules"):
            for item in chosen[field]:
                item["decision_status"] = "confirmed"
        saved_boundary = boundary.save_analysis_boundary(import_id, {
            "base_structure_revision_id": built["structure_revision_id"],
            "base_evidence_revision_id": built["evidence_revision_id"],
            "base_boundary_revision_id": None, "base_coverage_revision_id": None,
            **{key: chosen[key] for key in ("evaluation_objects", "source_scope_rules", "label_scope_rules")}}, LOGIN)
        self.assertTrue(saved_boundary["confirmation_ready"], saved_boundary)
        confirmed_boundary = boundary.confirm_analysis_boundary(import_id, {
            key: saved_boundary[key] for key in ("boundary_revision_id", "coverage_revision_id", "boundary_payload_sha256", "coverage_payload_sha256")}, LOGIN)
        from app.services import interview_v2_dossier_service as dossiers
        from app.services import interview_v2_analysis_service as analysis
        from app.services import interview_v2_report_service as reports
        from app.services import interview_v2_report_review_service as reviews
        from app.core.interview_v2_report import REPORT_SECTION_SPECS
        public, evidence, confirmed_rules, coverage, provenance = dossiers._ready_project(project_id, LOGIN)
        evidence_payload = evidence["evidence"]
        participants = evidence_payload["expected_participants"]
        self.assertEqual(len(participants), 2)
        recorder = llm_usage.start_llm_usage_task(LOGIN, category="interview", action="V2 连续流程离线验收")
        with bind_llm_api_key(FIXTURE_KEY), bind_llm_attempt_observer(recorder.on_attempt_event):
            for participant in participants:
                pid = participant["participant_id"]
                entry = next(item for item in evidence_payload["entries"] if item["participant_id"] == pid and item["inclusion_status"] == "included")
                self.answers["interview_v2.attribute"] = {"participant_id": pid, "facts": [], "analytical_labels": []}
                self.answers["interview_v2.dossier"] = {"participant_id": pid, "claims": [{"claim_type": "behavior",
                    "statement": entry["normalized_content"], "supporting_evidence_ids": [entry["evidence_id"]],
                    "conflicting_evidence_ids": []}], "contradictions": [], "missing_context": []}
                dossier = await dossiers.regenerate_dossier(project_id, pid, {"base_dossier_version_id": None}, LOGIN)
                approved = dossiers.review_dossier(project_id, pid, {"base_dossier_version_id": dossier["dossier_version_id"],
                    "decision": "approved", "note": "合成样本核对"}, LOGIN)
                self.assertEqual(approved["status"], "approved")
                self.assertEqual(store.load_current_participant_dossier(project_id, pid)["revision"]["llm_model_snapshot"]["version"], 1)
            evaluation = confirmed_rules["evaluation_objects"][0]
            cases = [{"participant_id": participant["participant_id"], "evidence_ids": [next(
                item["evidence_id"] for item in evidence_payload["entries"] if item["participant_id"] == participant["participant_id"]
                and item["inclusion_status"] == "included")]} for participant in participants]
            self.answers["interview_v2.analysis"] = {"module_id": evaluation["module_id"], "findings": [{
                "title": "组队方式", "statement": "玩家分别邀请好友和随机匹配。", "evaluation_object_id": evaluation["evaluation_object_id"],
                "main_question_id": evaluation["main_question_ids"][0], "supporting_cases": cases,
                "counterexample_cases": [], "observation_cases": [], "limitations": [], "confidence": 0.8}]}
            analyzed = await analysis.create_analysis_run(project_id, {"base_analysis_run_id": None, "freeze_current": True}, LOGIN)
            self.assertEqual(analyzed["status"], "completed")
            frozen_analysis = store.load_current_analysis_run(project_id)["revision"]
            finding_id = frozen_analysis["findings"][0]["finding_id"]
            types = {"scope_and_sample": "scope", "core_findings": "finding", "module_findings": "finding",
                     "participant_differences": "difference", "participant_logics": "logic", "recommendations": "suggestion", "evidence_and_limitations": "limitation"}
            sections = []
            for key, _title in REPORT_SECTION_SPECS:
                content = "玩家分别邀请好友和随机匹配。" if key == "core_findings" else "本章暂无额外发现。"
                sections.append({"section_key": key, "content": content, "claims": [{"claim_type": types[key], "text": content,
                    "start": 0, "end": len(content), "finding_ids": [] if key in {"scope_and_sample", "evidence_and_limitations"} else [finding_id],
                    "stat_fact_id": None}]})
            self.answers["interview_v2.report"] = {"sections": sections}
            self.answers["interview_v2.audit"] = {"issues": []}
            report = await reports.create_report(project_id, {"base_report_version_id": None, "freeze_current": True}, LOGIN)
            approved_report = reviews.approve_report(report["report_version_id"], {"base_report_version_id": report["report_version_id"]}, LOGIN)
            self.assertEqual(approved_report["status"], "approved")
        recorder.finish("completed")
        saved_report = store.load_current_report_version(project_id)["revision"]
        self.assertEqual(saved_report["llm_model_snapshot"]["version"], 1)
        self.assertEqual(frozen_analysis["llm_model_snapshot"]["version"], 1)
        self.assertEqual(len(self.calls), 7)
        self.assert_usage(["interview_v2.attribute", "interview_v2.dossier", "interview_v2.analysis", "interview_v2.report", "interview_v2.audit"])

    async def test_v2_report_entry_saves_writer_and_audit_settings(self):
        from app.services import interview_v2_report_service as service
        from tests import test_interview_v2_report_service as fixture
        service.store._mapping_lock_path(fixture.PROJECT).parent.mkdir(parents=True)
        self.answers["interview_v2.report"] = fixture._writer_output()
        self.answers["interview_v2.audit"] = {"issues": []}
        current = fixture._analysis()
        current["revision"]["revision_payload_sha256"] = service.store._analysis_digest(current["revision"])
        analysis_dir = service.store._analysis_dir(fixture.PROJECT)
        service.store._atomic_write_json(analysis_dir / "state.json", current["state"])
        service.store._atomic_write_json(analysis_dir / "versions" / (fixture.ANALYSIS + ".json"), current["revision"])
        recorder = llm_usage.start_llm_usage_task(LOGIN, category="interview", action="V2 报告验收")
        with patch.object(service, "get_current_analysis", return_value={"status": "completed"}), \
             patch.object(service.store, "load_current_analysis_run", return_value=current), \
             patch.object(service.store, "load_project", return_value={"project_id": fixture.PROJECT, "research_focus": "入口"}), \
             patch.object(service.store, "_require_analysis_source_current_locked"), \
             patch.object(service, "_frozen_config", return_value={"research_focus": "入口"}), \
             bind_llm_api_key(FIXTURE_KEY), bind_llm_attempt_observer(recorder.on_attempt_event):
            result = await service.create_report(fixture.PROJECT,
                {"base_report_version_id": None, "freeze_current": True}, LOGIN)
        recorder.finish("completed")
        self.assertEqual(result["status"], "draft")
        saved = service.store.load_current_report_version(fixture.PROJECT)
        self.assertEqual(saved["revision"]["llm_model_snapshot"]["version"], 1)
        self.assert_usage(["interview_v2.report", "interview_v2.audit"])

    async def test_v2_analysis_entry_saves_its_stage_settings(self):
        from app.services import interview_v2_analysis_service as service
        from tests import test_interview_v2_analysis_service as fixture
        service.store._mapping_lock_path(fixture.PROJECT).parent.mkdir(parents=True)
        self.answers["interview_v2.analysis"] = {"module_id": fixture.MODULE, "findings": [{
            "title": "合成发现", "statement": "一名玩家支持，另一名玩家提供反例。",
            "evaluation_object_id": fixture.OBJECT, "main_question_id": fixture.QUESTION,
            "supporting_cases": [{"participant_id": fixture.P1, "evidence_ids": [fixture.EV1]}],
            "counterexample_cases": [{"participant_id": fixture.P2, "evidence_ids": [fixture.EV2]}],
            "observation_cases": [], "limitations": [], "confidence": 0.8}]}
        recorder = llm_usage.start_llm_usage_task(LOGIN, category="interview", action="V2 综合分析验收")
        with patch.object(service, "_ready_project", return_value=fixture.ready_payload()), \
             patch.object(service.store, "load_current_participant_dossier", side_effect=lambda project, participant:
                 fixture.current_dossier(participant, "approved" if participant == fixture.P1 else "generated")), \
             patch.object(service.store, "_require_analysis_source_current_locked"), \
             bind_llm_api_key(FIXTURE_KEY), bind_llm_attempt_observer(recorder.on_attempt_event):
            result = await service.create_analysis_run(fixture.PROJECT,
                {"base_analysis_run_id": None, "freeze_current": True}, LOGIN)
        recorder.finish("completed")
        self.assertEqual(result["status"], "completed")
        saved = service.store.load_current_analysis_run(fixture.PROJECT)
        self.assertEqual(saved["revision"]["llm_model_snapshot"]["version"], 1)
        self.assert_usage(["interview_v2.analysis"])

    async def test_insight_report_entry_saves_run_version_history_and_usage(self):
        from app.services import survey_service as service
        from app.storage.report_postwrite_runs import RunStore
        from tests.test_report_postwrite_flow import REPORT
        from tests.test_report_postwrite_audit import session, response
        self.answers["insight.writer"] = REPORT
        def audit_answer(request):
            payload = json.loads(request["input"][-1]["content"])
            if "blocks" in payload:
                return {"claims": [{"id": "C1", "original": REPORT.split("\n\n")[1], "block_ids": ["B0002"],
                    "kind": "inference", "review_refs": ["q1/r4"]}],
                    "coverage": [{"block_id": "B0001", "no_claim": True}, {"block_id": "B0002", "claim_ids": ["C1"]}], "tasks": []}
            return {"checks": [response(item) for item in payload["targets"]]}
        self.answers.update({"insight.registry": audit_answer, "insight.meaning": audit_answer})
        sid = self.session(**session())
        self.stack.enter_context(patch.object(service, "is_postwrite_report_enabled", return_value=True))
        self.stack.enter_context(patch.object(service, "PostwriteRunStore", side_effect=lambda: RunStore(self.folder / "runs")))
        self.stack.enter_context(patch.object(service, "_current_login", AsyncMock(return_value=None)))
        self.stack.enter_context(patch.object(service, "_assign_session_owner"))
        events = await self.consume(service.report_stream(sid, Request({"type": "http", "headers": []})), sid=sid)
        self.assertIn("report_done", [item["type"] for item in events])
        saved = sessions.get_session(sid)
        self.assertIn(REPORT, saved["report_md"])
        self.assertEqual(saved["report_versions"][-1]["llm_model_snapshot"]["version"], 1)
        self.assertEqual(history_store._load_history()[0]["llm_model_snapshot"]["version"], 1)
        self.assert_usage(["insight.writer", "insight.registry", "insight.meaning"])

    async def test_quick_entry_saves_checkpoint_with_published_settings(self):
        from app.services import report_quick_pipeline as service
        from tests.test_report_quick_pipeline import question, answer
        self.answers["quick.summary"] = lambda request: answer(json.loads(request["input"][-1]["content"]))
        checkpoint = self.folder / "quick-checkpoint.json"
        async def persist(value):
            checkpoint.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        recorder = llm_usage.start_llm_usage_task(LOGIN, category="survey", action="快速报告验收")
        with bind_llm_api_key(FIXTURE_KEY), bind_llm_attempt_observer(recorder.on_attempt_event):
            result = await service.run_quick_pipeline([question()], on_checkpoint=persist,
                collect=llm_client.collect_chat_completion)
        recorder.finish("completed")
        self.assertEqual(result["report_status"], "complete")
        self.assertEqual(json.loads(checkpoint.read_text(encoding="utf-8"))["llm_model_snapshot"]["version"], 1)
        self.assert_usage(["quick.summary"])

    async def test_standard_report_and_qa_preserve_published_parameters_and_history(self):
        from app.services import survey_service as service
        from tests.test_survey_report_versions import _base_session, _writer
        source = _base_session()
        source.pop("owner_key")
        sid = self.session(**source)
        outputs = iter([answer for answer, _ in _writer("离线验收报告").side_effect])
        self.answers["survey.report"] = lambda request: next(outputs)
        self.answers["survey.qa"] = "此为合成追问结果。"
        self.stack.enter_context(patch.object(service, "_current_login", AsyncMock(return_value=None)))
        self.stack.enter_context(patch.object(service, "_assign_session_owner"))
        self.stack.enter_context(patch.object(service, "_ensure_branch_rules", return_value=[]))
        self.stack.enter_context(patch.object(service, "audit_log", AsyncMock()))
        self.stack.enter_context(patch.object(service, "is_postwrite_report_enabled", return_value=False))
        events = await self.consume(service.report_stream(sid, Request({"type": "http", "headers": []})), sid=sid)
        self.assertIn("report_done", [event["type"] for event in events])
        saved = sessions.get_session(sid)
        self.assertEqual(saved["report_versions"][-1]["llm_model_snapshot"]["version"], 1)
        writer_calls = len(self.calls)
        self.assertGreater(writer_calls, 0)
        self.assert_usage(["survey.report"])
        self.calls.clear()
        qa = await self.consume(service.qa_stream(sid, "合成报告说明了什么？", object()), sid=sid)
        self.assertIn("qa_done", [event["type"] for event in qa])
        self.assertIn("此为合成追问结果", json.dumps(sessions.get_session(sid), ensure_ascii=False))
        self.assert_usage(["survey.qa"])

    async def test_quant_report_and_qa_use_quant_stages(self):
        from app.services import survey_service as service
        from tests.test_survey_report_versions import _base_session, _writer
        source = _base_session()
        source.pop("owner_key")
        source.update(mode="quant", analysis_mode="quantitative")
        sid = self.session(**source)
        outputs = iter([answer for answer, _ in _writer("定量离线报告").side_effect])
        self.answers["quant.report"] = lambda request: next(outputs)
        self.answers["quant.qa"] = "此为定量合成追问结果。"
        for key, value in (("_current_login", AsyncMock(return_value=None)),
                           ("_assign_session_owner", lambda *args: None), ("audit_log", AsyncMock()),
                           ("_ensure_branch_rules", lambda _: []),
                           ("is_postwrite_report_enabled", lambda: False)):
            self.stack.enter_context(patch.object(service, key, value))
        events = await self.consume(service.report_stream(sid, Request({"type": "http", "headers": []})), sid=sid)
        self.assertIn("report_done", [event["type"] for event in events])
        self.assertEqual(sessions.get_session(sid)["report_versions"][-1]["llm_model_snapshot"]["version"], 1)
        self.assert_usage(["quant.report"])
        self.calls.clear()
        events = await self.consume(service.qa_stream(sid, "定量合成问题", object()), sid=sid)
        self.assertIn("qa_done", [event["type"] for event in events])
        self.assert_usage(["quant.qa"])

    async def test_business_http_budget_stops_before_an_extra_backup_request(self):
        from app.services import survey_service as service
        from tests.test_survey_planner_direct_llm import _plan_json
        state = (await self.client.get("/api/model-settings")).json()
        cfg = deepcopy(state["configuration"])
        cfg["tasks"]["survey.plan"]["max_http_attempts"] = 2
        published = await self.client.post("/api/model-settings/publish", json={
            "expected_revision": state["revision"], "configuration": cfg})
        self.assertEqual(published.status_code, 200, published.text)
        self.failures["survey.plan"] = "all503"
        self.answers["survey.plan"] = _plan_json()
        sid = self.session(rows=[["段位", "主玩位置", "界面满意度"], ["金", "中", "4"]])
        with patch.object(service, "audit_log", AsyncMock()), \
             patch.object(service, "_ensure_branch_rules", return_value=[]):
            events = await self.consume(service.plan_stream(sid, object()), sid=sid)
        self.assertIn("error", [event["type"] for event in events])
        self.assertNotIn("plan_ready", [event["type"] for event in events])
        self.assertEqual(len(self.calls), 2)
        self.assertTrue(self.calls[0]["model"].endswith("primary"))
        self.assertTrue(self.calls[1]["model"].endswith("backup"))
        self.assert_usage(["survey.plan"], missing=2, version=2)

    async def test_business_primary_failure_and_timeout_use_backup_without_losing_input(self):
        from app.services import survey_service as service
        from tests.test_survey_planner_direct_llm import _plan_json
        self.answers["survey.plan"] = _plan_json()
        for failure in ("503", "timeout", "unknown"):
            with self.subTest(failure=failure):
                self.calls.clear()
                self.failures["survey.plan"] = failure
                sid = self.session(rows=[["段位", "主玩位置", "界面满意度"], [PRIVATE_MATERIAL, "中", "4"]])
                with patch.object(service, "audit_log", AsyncMock()), \
                     patch.object(service, "_ensure_branch_rules", return_value=[]):
                    events = await self.consume(service.plan_stream(sid, object()), sid=sid)
                self.assertIn("plan_ready", [item["type"] for item in events])
                expected = 1 if failure == "unknown" else 2
                self.assertEqual(len(self.calls), expected)
                if expected == 2:
                    self.assertTrue(self.calls[-1]["model"].endswith("backup"))
                    self.assertEqual(self.calls[0]["input"], self.calls[1]["input"])
                    self.assertEqual(self.calls[0]["instructions"], self.calls[1]["instructions"])
                self.assertIn(PRIVATE_MATERIAL, json.dumps(self.calls[0], ensure_ascii=False))
                record = self.assert_usage(["survey.plan"], missing=1)
                self.assertEqual(record["status"], "completed")


if __name__ == "__main__":
    unittest.main()
