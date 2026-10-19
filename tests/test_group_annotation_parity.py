"""Persistent adapters for the 6b4740a views. --serve starts an isolated synthetic review."""
import asyncio
from copy import deepcopy
from io import BytesIO
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch, AsyncMock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
if '--serve' in sys.argv:
    WORK = Path(os.environ.get('PARITY_REVIEW_DIR') or tempfile.mkdtemp(prefix='survey-annotation-parity-'))
    WORK.mkdir(parents=True, exist_ok=True)
    os.chdir(WORK)  # Startup session sweeping is relative to cwd; isolate it too.
    os.environ.update(DATA_DIR=str(WORK/'data'), FEISHU_LOGIN_REQUIRED='false',
        LLM_API_KEY='synthetic-only', LLM_API_BASE='http://127.0.0.1:9',
        FEISHU_APP_ID='', FEISHU_APP_SECRET='', FEISHU_ADMIN_EMAILS='owner@example.test',
        FEISHU_WIKI_AUTO_UPDATE_ENABLED='false', FEISHU_NAVIGATION_NOTIFICATIONS_ENABLED='false')

from fastapi import HTTPException
from openpyxl import load_workbook, Workbook
from test_group_tasks import WorkflowCase, LOGIN, snapshot, entry
from app.core import config
from app.schemas.group_tasks import (TaskRules, FeedbackMapping, AnnotationOptions, RunTask,
    ConfirmAI, AIReviewDecision, QualityOverride)
from app.services import group_tasks as tasks, group_task_runner as runner
from app.storage import group_workflow as store


def create_sample(name='横屏流程验收', ai=True, quality=True, count=False):
    task = tasks.create_task(TaskRules(name=name, payout_date='2026-10-16', base_reward=50,
        excellent_reward=100, ai_detect=ai, requires_count=count, min_count=3 if count else None), LOGIN)
    content = 'Discord User ID,What do you think of the feature?,Why?,Matches\np1,The button freezes after two taps.,Please increase the button size.,3\np2,I like the feature.,,3\np3,Synthetic AI answer.,Synthetic explanation.,3\np4,I tried twice.,It was okay.,2\n'
    task = tasks.upload_feedback(task['id'],task['version'],'【合成验收】完整文件名-玩家反馈-20261016.csv',content.encode(),LOGIN)
    return tasks.set_mapping(task['id'],task['version'],FeedbackMapping(id_col=0,open_text_cols=[1,2],count_col=3 if count else None),LOGIN,
        AnnotationOptions(ai_detect=ai,quality=quality,background='研究玩家操作体验'))


def ai_result(row, probability=5, columns=(1,2)):
    answers={f'col_{c}':str(row['cells'][c] or '').strip() if c < len(row['cells']) else '' for c in columns}
    return {'id':row['id'],'ai_prob':probability,'polish_prob':8,'reason':'合成识别依据，仅用于流程验收',
        'evidence':next((value for value in answers.values() if value),''),'counter_evidence':'保留玩家具体体验作为反向依据',
        'translations':{key:'合成回答译文，仅用于流程验收。' if value else '' for key,value in answers.items()}}


def quality_result(row, columns=(1,2)):
    label='优秀反馈' if row['player_id']=='p1' else '有效反馈'
    answers={f'col_{c}':str(row['cells'][c] or '').strip() if c < len(row['cells']) else '' for c in columns}
    return {'id':row['id'],'overall':label,'overall_reason':'合成整体判断，仅用于流程验收',
        'q_labels':{key:label if value else 'N/A' for key,value in answers.items()},
        'q_reasons':{key:'根据该题和回答判断' if value else '该题未作答，按 N/A 处理' for key,value in answers.items()},
        'q_evidence':answers,
        'translations':ai_result(row,columns=columns)['translations']}


def seed_results(task, ai=True, quality=False, candidate=True):
    def save(current):
        for row in current['feedback']:
            if row['gate']['status']=='fail':continue
            if ai:row.update(ai=ai_result(row,90 if candidate and row['player_id']=='p3' else 5,current['mapping']['open_text_cols']),ai_complete=True)
            if quality:row.update(quality=quality_result(row,current['mapping']['open_text_cols']),quality_complete=True)
    store.mutate('tasks',task['id'],save)
    return tasks.task_view(store.read('tasks',task['id']),LOGIN)


class AnnotationParityTests(WorkflowCase):
    def test_shifted_multicolumn_workbooks_review_restore_and_export(self):
        layouts=[(0,[2,4,6,8,10]),(3,[8,16,17,18,19,20,22,23,24]),(5,[0,1,2,3,4,6,7,9,11,13,15,17])]
        for id_col,columns in layouts:
            with self.subTest(id_col=id_col,columns=columns):
                width=max(id_col,*columns)+1
                headers=[f'元数据{c}' for c in range(width)]
                headers[id_col]='Discord User ID'
                for col in columns:headers[col]=f'反馈题{col+1}'
                book=Workbook();sheet=book.active;sheet.append(headers)
                for player in ('p1','p2','p3'):
                    row=['']*width;row[id_col]=player
                    for col in columns:row[col]=f'{player} answer for question {col+1}'
                    if player=='p2':row[columns[-1]]=''
                    sheet.append(row)
                buffer=BytesIO();book.save(buffer);book.close()
                task=tasks.create_task(TaskRules(name=f'{len(columns)}题异位列验收',payout_date='2026-10-22',base_reward=100,excellent_reward=50,ai_detect=True),LOGIN)
                task=tasks.upload_feedback(task['id'],task['version'],'多题列位置.xlsx',buffer.getvalue(),LOGIN)
                task=tasks.set_mapping(task['id'],task['version'],FeedbackMapping(id_col=id_col,open_text_cols=columns),LOGIN,AnnotationOptions(ai_detect=True,quality=True))
                async def model(current,row,checkpoint):
                    if current['run']['stage']=='ai':
                        checkpoint({'ai':ai_result(row,90 if row['player_id']=='p3' else 5,columns),'ai_complete':True})
                    else:
                        checkpoint({'quality':quality_result(row,columns),'quality_complete':True})
                async def execute(stage):
                    current=store.read('tasks',task['id'])
                    with patch.object(runner,'require_background_llm_api_key',return_value='synthetic'),patch.object(runner,'start_llm_usage_task',return_value=None),patch.object(runner,'annotate_group_feedback',model):
                        runner.start_run(task['id'],RunTask(expected_version=current['version'],stage=stage),LOGIN)
                        for _ in range(10):
                            if store.read('tasks',task['id'])['run']['status']!='queued':break
                            await runner._run_slice(task['id'],asyncio.Semaphore(2))
                asyncio.run(execute('ai'))
                task=tasks.task_view(store.read('tasks',task['id']),LOGIN)
                self.assertEqual(task['annotation']['phase'],'ai_review')
                task=tasks.review_ai(task['id'],task['annotation']['review_ids'][0],AIReviewDecision(expected_version=task['version'],is_ai=False),LOGIN)
                task=tasks.confirm_ai(task['id'],ConfirmAI(expected_version=task['version'],response_ids=[]),LOGIN)
                asyncio.run(execute('quality'))
                task=tasks.task_view(store.read('tasks',task['id']),LOGIN)
                self.assertTrue(task['annotation']['complete'])
                self.assertEqual(task['annotation']['quality_done'],3)
                keys={f'col_{c}' for c in columns}
                for row in task['feedback']:
                    for field in ('q_labels','q_reasons','q_evidence','translations'):
                        self.assertEqual(set(row['quality'][field]),keys)
                rid=task['feedback'][0]['id'];col=columns[-1]
                task=tasks.override_quality(task['id'],rid,QualityOverride(expected_version=task['version'],column_index=col,label='无效反馈',note='跨列人工改标验收'),LOGIN)
                restored=tasks.task_view(store.read('tasks',task['id']),LOGIN)
                self.assertEqual(restored['feedback'][0]['quality']['q_labels'][f'col_{col}'],'无效反馈')
                self.assertIn(f'col_{col}',restored['feedback'][0]['quality']['human_reviews'])
                blob,_=tasks.annotation_download(task['id'],LOGIN)
                book=load_workbook(BytesIO(blob));values=list(book.active.values);book.close()
                label_col=values[0].index(f'[{headers[col]}]质量标注')
                self.assertEqual(values[1][label_col],'无效反馈')
                self.assertEqual(values[2][label_col],'N/A')
                self.assertEqual(values[3][label_col],'有效反馈')
                self.assertEqual(values[1][values[0].index('Discord User ID')],'p1')
                task=tasks.override_quality(task['id'],rid,QualityOverride(expected_version=task['version'],column_index=col,restore=True,note='恢复原标签'),LOGIN)
                self.assertNotIn(f'col_{col}',task['feedback'][0]['quality']['human_reviews'])
                self.assertEqual(task['feedback'][0]['quality']['q_labels'][f'col_{col}'],'优秀反馈')

    def test_stale_done_flag_cannot_enable_results_rewards_review_or_download(self):
        task=seed_results(create_sample(ai=False),ai=False,quality=True)
        rid=task['feedback'][0]['id']
        store.mutate('tasks',task['id'],lambda current:current.update(run={
            'stage':'quality','status':'completed',
            'completed_ids':[row['id'] for row in current['feedback']]}))
        for bad_label in (None,'未知标签','N/A'):
            with self.subTest(label=bad_label):
                def corrupt(current):
                    labels=current['feedback'][0]['quality']['q_labels']
                    if bad_label is None:labels.pop('col_1',None)
                    else:labels['col_1']=bad_label
                store.mutate('tasks',task['id'],corrupt)
                before=store.read('tasks',task['id'])
                view=tasks.task_view(before,LOGIN)
                self.assertEqual(view['annotation']['phase'],'quality_progress')
                self.assertEqual(view['annotation']['quality_done'],3)
                self.assertFalse(view['annotation']['complete'])
                self.assertFalse(view['reward_ready'])
                self.assertEqual(view['status'],'incomplete')
                self.assertNotIn(rid,view['run']['completed_ids'])
                self.assertFalse(tasks.stage_complete(before,before['feedback'][0],'quality'))
                with self.assertRaises(HTTPException):tasks.annotation_download(task['id'],LOGIN)
                with self.assertRaises(HTTPException):tasks.override_quality(task['id'],rid,QualityOverride(expected_version=task['version'],column_index=1,label='有效反馈',note='缺失标签不可人工绕过'),LOGIN)
                self.assertEqual(store.read('tasks',task['id']),before,'read/rejected actions must not repair stored data')
        with patch.object(runner,'require_background_llm_api_key',return_value='synthetic'):
            retry=runner.start_run(task['id'],RunTask(expected_version=task['version'],stage='quality'),LOGIN)
        self.assertEqual(retry['run']['target_ids'],[rid])

    def test_real_adapter_rechecks_labels_despite_cached_done_flag(self):
        from app.services import annotate_workflow as workflow
        task=seed_results(create_sample(ai=False),ai=False,quality=True)
        row=deepcopy(task['feedback'][0]);row['quality']['q_labels'].pop('col_1')
        task['run']={'stage':'quality'}
        repaired=quality_result(row,task['mapping']['open_text_cols'])
        strict=AsyncMock(return_value=(1,[repaired],set(),''));updates=[]
        with patch.object(workflow,'_run_one_quality_batch_strict',strict):
            asyncio.run(workflow.annotate_group_feedback(task,row,updates.append))
        strict.assert_awaited_once()
        self.assertTrue(updates[-1]['quality_complete'])
        self.assertEqual(updates[-1]['quality']['q_labels']['col_1'],'优秀反馈')

    def test_none_selected_rejected_without_saving_mapping_or_options(self):
        task=create_sample()
        with self.assertRaises(HTTPException):
            tasks.set_mapping(task['id'],task['version'],FeedbackMapping(**task['mapping']),LOGIN,AnnotationOptions(ai_detect=False,quality=False))
        self.assertEqual(store.read('tasks',task['id'])['version'],task['version'])

    def test_options_persist_and_identical_confirmation_keeps_results(self):
        task=seed_results(create_sample(),quality=True)
        row=task['feedback'][0]
        updated=tasks.set_mapping(task['id'],task['version'],FeedbackMapping(**task['mapping']),LOGIN,AnnotationOptions(**task['annotation']['options']))
        self.assertEqual(updated['feedback'][0]['ai'],row['ai'])
        self.assertEqual(updated['annotation']['options']['background'],'研究玩家操作体验')

    def test_context_change_requires_new_task_and_keeps_results(self):
        task=seed_results(create_sample(),quality=True)
        before=store.read('tasks',task['id'])
        with self.assertRaises(HTTPException):
            tasks.set_mapping(task['id'],task['version'],FeedbackMapping(**task['mapping']),LOGIN,AnnotationOptions(background='另一个研究背景'))
        self.assertEqual(store.read('tasks',task['id']),before)

    def test_metadata_settings_keep_ai_confirmation_and_question_review(self):
        from app.schemas.group_tasks import TaskSettingsUpdate
        task=seed_results(create_sample(),quality=True,candidate=False)
        task=tasks.confirm_ai(task['id'],ConfirmAI(expected_version=task['version'],response_ids=[]),LOGIN)
        task=tasks.override_quality(task['id'],task['feedback'][0]['id'],QualityOverride(expected_version=task['version'],column_index=1,label='有效反馈',note='保留人工判定'),LOGIN)
        before=deepcopy(task['feedback'])
        updated=tasks.update_settings(task['id'],TaskSettingsUpdate(expected_version=task['version'],name='改名称和日期',payout_date='2026-10-22',base_reward=200,excellent_reward=100),LOGIN)
        self.assertTrue(updated['ai_confirmation_complete'])
        self.assertTrue(updated['reward_ready'])
        self.assertEqual(updated['annotation']['phase'],'results')
        for old,new in zip(before,updated['feedback']):
            for key in ('ai','quality','reviews','ai_reviewed','confirmed_ai'):
                self.assertEqual(new.get(key),old.get(key),key)

    def test_progress_partial_results_cannot_enter_review_or_download(self):
        task=create_sample()
        store.mutate('tasks',task['id'],lambda t:t['feedback'][0].update(ai=ai_result(t['feedback'][0]),ai_complete=True))
        task=tasks.task_view(store.read('tasks',task['id']),LOGIN)
        self.assertEqual((task['annotation']['phase'],task['annotation']['ai_done']),('ai_progress',1))
        with self.assertRaises(HTTPException):tasks.annotation_download(task['id'],LOGIN)
        with self.assertRaises(HTTPException):tasks.confirm_ai(task['id'],ConfirmAI(expected_version=task['version'],response_ids=[]),LOGIN)

    def test_candidates_require_individual_review_and_survive_reload(self):
        task=seed_results(create_sample())
        self.assertEqual(task['annotation']['phase'],'ai_review')
        rid=task['annotation']['review_ids'][0]
        with self.assertRaises(HTTPException):tasks.confirm_ai(task['id'],ConfirmAI(expected_version=task['version'],response_ids=[rid]),LOGIN)
        task=tasks.review_ai(task['id'],rid,AIReviewDecision(expected_version=task['version'],is_ai=True),LOGIN)
        restored=tasks.task_view(store.read('tasks',task['id']),LOGIN)
        self.assertEqual(restored['annotation']['reviewed_ids'],[rid])
        task=tasks.confirm_ai(task['id'],ConfirmAI(expected_version=task['version'],response_ids=[rid]),LOGIN)
        self.assertEqual(task['annotation']['phase'],'quality_progress')
        self.assertEqual((task['annotation']['quality_total'],task['annotation']['quality_done']),(3,0))
        self.assertEqual(task['feedback'][2]['outcome']['source'],'human_ai')

    def test_ai_only_finishes_and_exports_without_inventing_quality_or_rewards(self):
        task=seed_results(create_sample(quality=False),candidate=False)
        task=tasks.confirm_ai(task['id'],ConfirmAI(expected_version=task['version'],response_ids=[]),LOGIN)
        self.assertTrue(task['annotation']['complete'])
        self.assertEqual(task['status'],'annotation_done')
        self.assertFalse(task['workflow']['annotation_complete'])
        with self.assertRaises(HTTPException):tasks.confirm_task(task['id'],task['version'],LOGIN)
        blob,name=tasks.annotation_download(task['id'],LOGIN)
        book=load_workbook(BytesIO(blob));headers=[c.value for c in book.active[1]]
        self.assertIn('AI内容生成概率',headers);self.assertNotIn('整体反馈质量',headers)
        self.assertNotIn('研究玩家操作体验',str(list(book.active.values)))
        self.assertEqual(book.active.max_row,5);book.close()

    def test_quality_only_skips_ai_and_exports_original_player_identity(self):
        task=seed_results(create_sample(ai=False),ai=False,quality=True)
        self.assertEqual(task['annotation']['phase'],'results')
        blob,_=tasks.annotation_download(task['id'],LOGIN)
        book=load_workbook(BytesIO(blob));values=list(book.active.values);headers=values[0]
        self.assertNotIn('AI内容生成概率',headers)
        self.assertIn('整体反馈质量',headers)
        self.assertEqual(values[1][headers.index('Discord User ID')],'p1')
        self.assertNotIn(task['feedback'][0]['id'],str(values));book.close()

    def test_question_review_restores_original_and_cannot_edit_blank_or_confirmed_ai(self):
        task=seed_results(create_sample(ai=False),ai=False,quality=True)
        rid=task['feedback'][0]['id']
        task=tasks.override_quality(task['id'],rid,QualityOverride(expected_version=task['version'],column_index=1,label='无效反馈',note='逐题复核'),LOGIN)
        self.assertEqual(task['feedback'][0]['quality']['q_labels']['col_1'],'无效反馈')
        self.assertIn('col_1',task['feedback'][0]['quality']['human_reviews'])
        task=tasks.override_quality(task['id'],rid,QualityOverride(expected_version=task['version'],column_index=1,restore=True,note='恢复'),LOGIN)
        self.assertEqual(task['feedback'][0]['quality']['q_labels']['col_1'],'优秀反馈')
        self.assertEqual(task['feedback'][0]['quality']['human_reviews'],{})
        with self.assertRaises(HTTPException):tasks.override_quality(task['id'],task['feedback'][1]['id'],QualityOverride(expected_version=task['version'],column_index=2,label='有效反馈',note='空题'),LOGIN)

    def test_another_operator_cannot_review_or_export(self):
        task=seed_results(create_sample(ai=False),ai=False,quality=True)
        with self.assertRaises(HTTPException):tasks.annotation_download(task['id'],{'email':'other@example.test'})
        with self.assertRaises(HTTPException):tasks.override_quality(task['id'],task['feedback'][0]['id'],QualityOverride(expected_version=task['version'],column_index=1,label='无效反馈',note='复核'),{'email':'other@example.test'})

    def test_question_review_keeps_independent_holistic_overall(self):
        import annotate
        task=seed_results(create_sample(ai=False),ai=False,quality=True)
        store.mutate('tasks',task['id'],lambda t:t['feedback'][0]['quality'].update(quality_policy_version=annotate.QUALITY_POLICY_VERSION))
        task=tasks.override_quality(task['id'],task['feedback'][0]['id'],QualityOverride(expected_version=task['version'],column_index=1,label='无效反馈',note='逐题复核'),LOGIN)
        result=task['feedback'][0]['quality']
        self.assertEqual(result['q_labels']['col_1'],'无效反馈')
        self.assertEqual(result['overall'],'优秀反馈')
        self.assertIn('未重新评估',result['overall_review_note'])

    def test_count_gate_export_preserves_original_without_pending_or_invented_ai_scores(self):
        task=seed_results(create_sample(count=True),quality=True,candidate=False)
        task=tasks.confirm_ai(task['id'],ConfirmAI(expected_version=task['version'],response_ids=[]),LOGIN)
        blob,_=tasks.annotation_download(task['id'],LOGIN)
        book=load_workbook(BytesIO(blob));values=list(book.active.values);headers=values[0]
        row=values[4]
        self.assertEqual(row[headers.index('Discord User ID')],'p4')
        self.assertEqual(row[headers.index('整体反馈质量')],'无效反馈')
        self.assertIn('数据校验未通过',row[headers.index('AI作答标签')])
        self.assertIsNone(row[headers.index('AI内容生成概率')])
        self.assertNotIn('待补齐',str(row));book.close()

    def test_unselected_quality_stage_is_rejected(self):
        task=seed_results(create_sample(quality=False),candidate=False)
        task=tasks.confirm_ai(task['id'],ConfirmAI(expected_version=task['version'],response_ids=[]),LOGIN)
        with self.assertRaises(HTTPException):runner.start_run(task['id'],RunTask(expected_version=task['version'],stage='quality'),LOGIN)

    def test_no_candidate_auto_continues_quality_without_browser(self):
        task=create_sample()
        async def model(current,row,checkpoint):
            if current['run']['stage']=='ai':checkpoint({'ai':ai_result(row),'ai_complete':True})
            else:checkpoint({'quality':quality_result(row),'quality_complete':True})
        async def execute():
            with patch.object(runner,'require_background_llm_api_key',return_value='synthetic'),patch.object(runner,'start_llm_usage_task',return_value=None),patch.object(runner,'annotate_group_feedback',model):
                runner.start_run(task['id'],RunTask(expected_version=task['version']),LOGIN)
                for _ in range(12):
                    current=store.read('tasks',task['id'])
                    if current['run']['status']!='queued':break
                    await runner._run_slice(task['id'],asyncio.Semaphore(2))
        asyncio.run(execute())
        restored=tasks.task_view(store.read('tasks',task['id']),LOGIN)
        self.assertEqual(restored['annotation']['phase'],'results')
        self.assertTrue(restored['ai_confirmation_complete'])
        self.assertEqual(restored['run']['stage'],'quality')

    def test_retry_only_missing_and_preserves_successful_ai_results(self):
        task=create_sample(quality=False)
        first=task['feedback'][0]['id']
        store.mutate('tasks',task['id'],lambda t:t['feedback'][0].update(ai=ai_result(t['feedback'][0]),ai_complete=True))
        with patch.object(runner,'require_background_llm_api_key',return_value='synthetic'):
            task=runner.start_run(task['id'],RunTask(expected_version=task['version'],stage='ai'),LOGIN)
        self.assertNotIn(first,task['run']['target_ids'])
        self.assertEqual(len(task['run']['target_ids']),3)

    def test_background_is_only_sent_to_ai_adapter(self):
        from app.services import annotate_workflow as workflow
        task=create_sample()
        row=task['feedback'][0]
        task['run']={'stage':'ai'}
        ai=AsyncMock(return_value=(1,[ai_result(row)],set(),set(),''))
        with patch.object(workflow,'_run_ai_batch_checked',ai):
            asyncio.run(workflow.annotate_group_feedback(task,row,lambda update:None))
        self.assertEqual(ai.await_args.args[-1],'研究玩家操作体验')
        task['run']={'stage':'quality'};task['ai_confirmation_complete']=True
        row.update(ai=ai_result(row),ai_complete=True)
        quality=AsyncMock(return_value=(1,[quality_result(row)],set(),''))
        with patch.object(workflow,'_run_one_quality_batch_strict',quality):
            asyncio.run(workflow.annotate_group_feedback(task,row,lambda update:None))
        self.assertNotEqual(quality.await_args.kwargs['background'],'研究玩家操作体验')


def serve_review():
    """Actual HTTP/UI + durable runner. Model and Feishu boundaries are deterministic doubles."""
    from app.services import auth, annotate_workflow, roster_sources
    from app.schemas.roster_sources import RosterSourceInput
    from app.integrations import feishu_sheets
    LOGIN['token']='synthetic-only'
    async def login(request):return LOGIN
    auth._current_login=login
    empty_start = '--empty' in sys.argv
    seed_sources = [] if empty_start else [('actual_group','实际名单群'),('note_group','任务备注群')]
    for identifier,name in seed_sources:
        if not store.read('sources',identifier):
            source={**RosterSourceInput(name=name,spreadsheet_url='syntheticbook123456',sheet_id=identifier,player_col=0,ori_col=1,status_col=2).model_dump(),
                'id':identifier,'revision':1,'sheet_title':name,'spreadsheet':'syntheticbook123456',
                'last_read_at':None,'last_read_status':'尚未读取','row_count':4 if identifier=='actual_group' else 0}
            store.create('sources',identifier,source)
    async def read_sheet(url,sheet_id,token,**kwargs):
        if url!='syntheticbook123456' or sheet_id not in ('actual_group','note_group'):raise feishu_sheets.SheetsError('隔离验收仅支持合成名单')
        rows=[[f'p{i}',f'{800000000+i}(6001)','群内人员'] for i in range(1,5)] if sheet_id=='actual_group' else []
        return {'spreadsheet':url,'sheet':{'sheet_id':sheet_id,'title':'实际名单群'},'values':[['playerid','ori','当前状态'],*rows]}
    feishu_sheets.read_sheet=read_sheet
    async def list_sheets(spreadsheet,token):
        if spreadsheet!='syntheticbook123456':raise feishu_sheets.SheetsError('隔离验收仅支持合成名单')
        return [{'sheet_id':sid,'title':name,'grid_properties':{'column_count':3,'row_count':5}}
                for sid,name in [('actual_group','实际名单群'),('note_group','任务备注群')]]
    async def deny_external(*args,**kwargs):raise feishu_sheets.SheetsError('隔离验收禁止外部飞书请求')
    feishu_sheets.list_sheets=list_sheets
    feishu_sheets._get=deny_external
    async def translate(headers):
        zh={'Discord User ID':'Discord 用户 ID','What do you think of the feature?':'你对这个功能有什么看法？','Why?':'为什么？','Matches':'实际体验次数'}
        return [zh.get(h,h) for h in headers],''
    annotate_workflow._translate_headers=translate
    async def model(task,row,checkpoint):
        await asyncio.sleep(1.2)
        if '失败重试' in task['rules']['name'] and row['player_id']=='p2' and not row.get('model_error'):
            checkpoint({'model_error':'合成失败：模型结果未通过完整性校验，请重试缺失反馈。'})
            return
        columns=task['mapping']['open_text_cols']
        if task['run']['stage']=='ai':checkpoint({'ai':ai_result(row,90 if row['player_id']=='p3' else 5,columns),'ai_complete':True,'model_error':''})
        else:checkpoint({'quality':quality_result(row,columns),'quality_complete':True,'model_error':''})
    runner.annotate_group_feedback=model
    runner.require_background_llm_api_key=lambda login:'synthetic-only'
    runner.start_llm_usage_task=lambda *args,**kwargs:None
    async def formula(source,login):return {'target_spreadsheet':'syntheticbook123456','sheet_title':'实际名单群','player_col':0,'ori_col':1,'status_col':2}
    roster_sources.validate_formula_reference=formula
    if not empty_start and not store.list_records('tasks'):
        for name,mode in [('列确认与任务选择','columns'),('AI 人工确认','review'),('质量结果预览','results')]:
            task=create_sample('合成验收：'+name,ai=mode!='results')
            def metadata(current):current.update(headers_zh=['Discord 用户 ID','你对这个功能有什么看法？','为什么？','实际体验次数'],header_translation_warning='')
            store.mutate('tasks',task['id'],metadata)
            if mode=='review':seed_results(task)
            if mode=='results':seed_results(task,ai=False,quality=True)
    book=Workbook();sheet=book.active
    sheet.append(['Discord User ID','What do you think of the feature?','Why?','Matches'])
    for row in [('p1','The button freezes after two taps.','Please increase the button size.',3),('p2','I like the feature.','',3),('p3','Synthetic AI answer.','Synthetic explanation.',3),('p4','I tried twice.','It was okay.',2)]:sheet.append(row)
    fixture=WORK/'【合成验收】完整文件名-四份玩家反馈-20261016.xlsx';book.save(fixture);book.close()
    from app.main import app
    from app.routers import group_tasks as routes
    async def key(request):return 'synthetic-only'
    routes.require_request_llm_api_key=key
    async def announcement(text, api_key):
        await asyncio.sleep(0.6)
        if '失败' in text:
            raise HTTPException(502,'合成识别失败，请重试')
        return {'suggestions':{
            'name':{'value':'合成验收：完整流程失败重试','evidence':text},
            'base_reward':{'value':100,'evidence':'valid answers grants 100 diamonds'},
            'excellent_reward':{'value':200,'evidence':'additional 200 diamonds'}},'warnings':[]}
    routes.extract_announcement=announcement
    print('SYNTHETIC_REVIEW_DIR='+str(WORK),flush=True)
    print('SYNTHETIC_UPLOAD='+str(fixture),flush=True)
    import uvicorn
    uvicorn.run(app,host='127.0.0.1',port=int(os.environ.get('PARITY_PORT','8773')),access_log=False)


if __name__=='__main__':
    if '--serve' in sys.argv:serve_review()
    else:unittest.main()
