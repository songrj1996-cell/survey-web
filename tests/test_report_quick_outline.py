"""Pure outline contracts: no model calls, storage writes or reordered answers."""
from copy import deepcopy
import re
import unittest

from app.services.report_modes import prepare_report_markdown
from app.services.report_quick_mode import render_quick_report
from app.services.report_quick_outline import build_quick_outline, outline_quick_markdown


def outline_fixture():
    stats = [{"question_key": str(i), "source_order": i + 1, "question": name,
              "markdown": f"### {name}\n\n| 选项 | 频数 |\n|---|---|\n| 合成值 | 3 |"}
             for i, name in enumerate(("打野经历", "性别"))]
    questions = [{"question_key": str(i), "source_order": i + 1, "question": name,
                  "status": "failed" if i == 5 else "complete",
                  "sources": [{"response_id": f"{i}/r1", "text": f"合成原文{i}", "profile": {"段位": "Gold"}}],
                  "findings": [] if i == 5 else [{"text": f"观点{i}", "frequency": "部分提及", "risk": False, "evidence_ids": [f"{i}/r1"]}]}
                 for i, name in enumerate(("共同认知", "常用英雄", "体验", "常用英雄"), 2)]
    def rule(option, indexes):
        return {"parent_index": 0, "parent_name": "打野经历", "allowed_options": [option],
                "confidence": "high", "source": "inferred_from_responses",
                "targets": [{"indexes": [i], "name": questions[i-2]["question"]} for i in indexes]}
    summary = {"questions": questions, "objective_stats": {"sections": stats}, "report_status": "partial"}
    snapshot = {"report_mode": "quick", "report_style": "quick", "report_status": "partial",
                "quick_summary": summary, "title": "合成报告",
                "input_snapshot": {"confirmed_columns": [{"column_indexes": [i]} for i in range(6)],
                                   "branch_rules": [rule("能玩但非主选", [3, 4]), rule("当前主玩", [5])]},
                "quick_checkpoint": {"untouched": True}, "qa_messages": [{"role": "user", "content": "保留追问"}]}
    snapshot["report_md"] = render_quick_report(summary, title=snapshot["title"])
    return snapshot


class QuickOutlineTests(unittest.TestCase):
    def test_actual_docx_keeps_group_and_question_heading_levels(self):
        from io import BytesIO
        from docx import Document
        from app.services.report_render import markdown_to_docx
        snapshot = outline_fixture()
        document = Document(BytesIO(markdown_to_docx(prepare_report_markdown(snapshot))))
        self.assertEqual([p.text for p in document.paragraphs if p.style.name == 'Heading 2'],
                         ['基础信息与共同问题', '能玩但非主选', '当前主玩'])
        self.assertEqual([p.text for p in document.paragraphs if p.style.name == 'Heading 3'],
                         ['打野经历', '性别', '共同认知', '常用英雄', '体验', '常用英雄'])
        self.assertEqual(len(document.tables), 2)
        self.assertIn('本题尚未完成，请重试。', [p.text for p in document.paragraphs])

    def test_all_leaves_once_with_duplicate_titles_and_no_snapshot_mutation(self):
        snapshot = outline_fixture()
        before = deepcopy(snapshot)
        outline = build_quick_outline(snapshot)
        self.assertEqual(outline["question_keys"], [str(i) for i in range(6)])
        self.assertEqual([g["title"] for g in outline["groups"]], ["基础信息与共同问题", "能玩但非主选", "当前主玩"])
        self.assertEqual([key for g in outline["groups"] for key in g["question_keys"]], outline["question_keys"])
        self.assertTrue(all("推定" in g["note"] for g in outline["groups"][1:]))
        displayed = prepare_report_markdown(snapshot)
        self.assertEqual(re.findall(r"^### (.*)$", displayed, re.M), ["打野经历", "性别", "共同认知", "常用英雄", "体验", "常用英雄"])
        # Each original question body survives byte-for-byte; failed placeholders
        # and objective tables are as important as completed findings.
        original_sections = re.split(r"^## [^\n]+\n", snapshot["report_md"], flags=re.M)[1:]
        displayed_sections = re.split(r"^### [^\n]+\n", displayed, flags=re.M)[1:]
        for original, shown in zip(original_sections, displayed_sections):
            self.assertTrue(shown.startswith(original), (original, shown))
        self.assertEqual(snapshot, before)
        self.assertTrue(prepare_report_markdown(snapshot, "evidence").startswith(displayed))

    def test_missing_legacy_malformed_and_edited_alignment_keep_flat_content(self):
        changes = [lambda s: s.update(report_mode="insight"),
                   lambda s: s.pop("input_snapshot"),
                   lambda s: s["input_snapshot"].update(branch_rules=[]),
                   lambda s: s["input_snapshot"].update(branch_rules=[None]),
                   lambda s: s["input_snapshot"]["branch_rules"][0].update(targets="bad"),
                   lambda s: s["quick_summary"]["questions"][1].update(question="edited"),
                   lambda s: s["quick_summary"]["questions"][1].update(question_key="2"),
                   lambda s: s.update(report_md=s["report_md"] + "\n## 无关附录\n"),
                   lambda s: s.update(report_md='```\n' + s['report_md'] + '\n```'),
                   lambda s: s["quick_summary"].update(objective_stats="bad")]
        for change in changes:
            with self.subTest(change=change):
                snapshot = outline_fixture()
                change(snapshot)
                before = deepcopy(snapshot)
                self.assertIsNone(build_quick_outline(snapshot))
                self.assertEqual(outline_quick_markdown(snapshot, snapshot["report_md"]), snapshot["report_md"])
                self.assertEqual(snapshot, before)

    def test_conflicting_medium_and_nested_rules_leave_affected_questions_independent(self):
        for kind in ("conflict", "medium", "nested", "partial_index"):
            snapshot = outline_fixture()
            rules = snapshot["input_snapshot"]["branch_rules"]
            if kind == "conflict":
                rules.append({**deepcopy(rules[0]), "allowed_options": ["冲突条件"]})
            elif kind == "medium":
                rules[0]["confidence"] = "medium"
            elif kind == "nested":
                rules[0]["parent_index"] = 5
            else:
                rules[0]["targets"][0]["indexes"] = [3, 99]
                rules[0]["targets"][1]["indexes"] = [4, 99]
            outline = build_quick_outline(snapshot)
            self.assertIsNotNone(outline)
            group = next(g for g in outline["groups"] if "3" in g["question_keys"])
            self.assertEqual(group["title"], "", kind)
            self.assertIn("## 常用英雄\n", prepare_report_markdown(snapshot))

    def test_noncontiguous_groups_preserve_order_and_repeat_as_continuation(self):
        snapshot = outline_fixture()
        rules = snapshot["input_snapshot"]["branch_rules"]
        rules[0]["targets"] = [{"indexes": [3]}, {"indexes": [5]}]
        rules[1]["targets"] = [{"indexes": [4]}]
        outline = build_quick_outline(snapshot)
        self.assertEqual([g["title"] for g in outline["groups"]],
                         ["基础信息与共同问题", "能玩但非主选", "当前主玩", "能玩但非主选（续）"])
        self.assertEqual([key for g in outline["groups"] for key in g["question_keys"]], [str(i) for i in range(6)])

    def test_conditions_are_notes_and_untrusted_group_labels_are_escaped(self):
        snapshot = outline_fixture()
        question = snapshot["quick_summary"]["questions"][1]
        question["question"] += "【推定适用于「打野经历」选择「能玩但非主选」的玩家】"
        snapshot["report_md"] = render_quick_report(snapshot["quick_summary"], title=snapshot["title"])
        snapshot["input_snapshot"]["branch_rules"][0]["allowed_options"] = ["<img src=x>\n## 伪标题"]
        displayed = prepare_report_markdown(snapshot)
        self.assertNotIn("<img", displayed)
        self.assertNotIn("\n## 伪标题", displayed)
        self.assertIn("### 常用英雄\n\n推定适用于", displayed)

    def test_mismatched_export_markdown_is_not_projected(self):
        snapshot = outline_fixture()
        other = snapshot["report_md"].replace("## 体验", "## 改过的标题")
        self.assertEqual(outline_quick_markdown(snapshot, other), other)


if __name__ == "__main__":
    unittest.main()
